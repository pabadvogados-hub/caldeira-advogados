"""
Contrato, procuracao e declaracao preenchidos a partir da qualificacao e da triagem.
Os modelos ficam em DOCS_MODELOS/ com {{CAMPOS}} (lista em docs/CAMPOS_DOS_MODELOS.md).
"""
import json
import os
import re
import sys
from datetime import date, datetime, timedelta

from docx import Document

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from docx_caldeira import preencher_modelo  # noqa: E402
from pasta_cliente import CONTRATACAO, nome_pasta, raiz_clientes  # noqa: E402
from relatorio_triagem import endereco_completo  # noqa: E402

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, RAIZ)
from config.escritorio import ESCRITORIO, OUTORGADOS, PRAZOS, TITULAR  # noqa: E402

MODELOS = os.path.join(RAIZ, 'DOCS_MODELOS')
DOCUMENTOS = [
    ('Contrato de Honorarios', 'CONTRATO_HONORARIOS_MODELO.docx'),
    ('Procuracao', 'PROCURACAO_MODELO.docx'),
    ('Declaracao de Hipossuficiencia', 'DECLARACAO_HIPOSSUFICIENCIA_MODELO.docx'),
]
MESES = ['janeiro', 'fevereiro', 'marco', 'abril', 'maio', 'junho', 'julho', 'agosto', 'setembro',
         'outubro', 'novembro', 'dezembro']
PENDENCIA = re.compile(r'\[(?:PREENCHER|CONFERIR)[^\]]*\]')


def data_extenso(d=None):
    d = d or date.today()
    mes = MESES[d.month - 1].replace('marco', 'março')
    return f'{d.day} de {mes} de {d.year}'


def _data(txt):
    try:
        return datetime.strptime((txt or '').strip()[:10], '%d/%m/%Y').date()
    except ValueError:
        return None


def data_contrato(cadastro, triagem):
    return (_data(cadastro.get('data_contrato') or cadastro.get('data_do_contrato'))
            or _data(triagem.get('data_reuniao')) or date.today())


def prazos_do_caso(inicio):
    """Marcos da fase de contratacao ate o protocolo da inicial."""
    p = PRAZOS
    onboarding = inicio + timedelta(days=p['onboarding_dias_apos_contrato'])
    fmt = lambda d: d.strftime('%d/%m/%Y')  # noqa: E731
    return [
        {'marco': 'Pedir documentos do checklist ao cliente', 'id': 'documentos',
         'data': fmt(inicio + timedelta(days=p['documentos_dias_apos_contrato'])), 'responsavel': 'Estagiário'},
        {'marco': 'Reunião de onboarding com o cliente', 'id': 'onboarding',
         'data': fmt(onboarding), 'responsavel': 'Gestor Jurídico'},
        {'marco': 'Notificação extrajudicial enviada aos bancos', 'id': 'notificacao',
         'data': fmt(onboarding + timedelta(days=p['notificacao_dias_apos_onboarding'])),
         'responsavel': 'Adv. Extrajudicial'},
        {'marco': 'Protocolo da petição inicial (prazo máximo)', 'id': 'inicial',
         'data': fmt(inicio + timedelta(days=p['inicial_dias_apos_contrato'])), 'responsavel': 'Adv. Judicial'},
    ]


def bancos(triagem):
    vistos = []
    for o in triagem.get('operacoes') or []:
        b = (o.get('banco') or '').strip()
        if b and b.lower() not in [v.lower() for v in vistos]:
            vistos.append(b)
    if not vistos:
        return ''
    return vistos[0] if len(vistos) == 1 else ', '.join(vistos[:-1]) + ' e ' + vistos[-1]


def proximo_numero_contrato():
    caminho = os.path.join(raiz_clientes(), '_controle_contratos.json')
    os.makedirs(os.path.dirname(caminho), exist_ok=True)
    try:
        with open(caminho, encoding='utf-8') as f:
            controle = json.load(f)
    except (FileNotFoundError, json.JSONDecodeError):
        controle = {}
    ano = str(date.today().year)
    controle[ano] = controle.get(ano, 0) + 1
    with open(caminho, 'w', encoding='utf-8') as f:
        json.dump(controle, f, indent=2)
    return f'{controle[ano]:03d}/{ano}'


def montar_campos(qualif, triagem, cadastro):
    h = triagem.get('honorarios') or {}
    lista_bancos = cadastro.get('bancos') or bancos(triagem)
    orgao = qualif.get('orgao_emissor', '')
    return {
        'NOME': (qualif.get('nome') or '').upper(),
        'NACIONALIDADE': qualif.get('nacionalidade', ''),
        'ESTADO_CIVIL': qualif.get('estado_civil', ''),
        'PROFISSAO': qualif.get('profissao', ''),
        'RG': qualif.get('rg', ''),
        'ORGAO_EMISSOR': orgao,
        'CPF': qualif.get('cpf', ''),
        'ENDERECO_COMPLETO': endereco_completo(qualif),
        'TELEFONE': qualif.get('telefone', ''),
        'EMAIL': qualif.get('email', ''),
        'BANCOS': lista_bancos,
        'OBJETO': cadastro.get('objeto') or (
            f'Defesa do(a) Outorgante nas operações de crédito rural mantidas com as seguintes instituições credoras: {lista_bancos}, '
            'nas esferas extrajudicial e judicial.' if lista_bancos else ''),
        'OUTORGADOS': '; '.join(f"{o['nome']}, {o['oab']}" for o in OUTORGADOS),
        'HONORARIOS_ENTRADA': cadastro.get('honorarios_entrada') or h.get('entrada', ''),
        'HONORARIOS_PAGAMENTO': cadastro.get('honorarios_pagamento') or h.get('parcelas', ''),
        'HONORARIOS_EXITO': cadastro.get('honorarios_exito') or h.get('exito', ''),
        'CIDADE_UF': f"{ESCRITORIO['cidade']}/{ESCRITORIO['uf']}",
        'DATA_EXTENSO': data_extenso(),
        'TITULAR_NOME': TITULAR['nome'].upper(),
        'TITULAR_OAB': TITULAR['oab'],
        'ESCRITORIO_CNPJ': ESCRITORIO['cnpj'],
        'ESCRITORIO_ENDERECO': ESCRITORIO['endereco'],
        'ESCRITORIO_EMAIL': ESCRITORIO['email'],
    }


def gerar(base, qualif, triagem, cadastro):
    campos = montar_campos(qualif, triagem, cadastro)
    campos['NUMERO_CONTRATO'] = proximo_numero_contrato()
    nome = nome_pasta(qualif.get('nome'))
    gerados = []
    for titulo_doc, modelo in DOCUMENTOS:
        caminho_modelo = os.path.join(MODELOS, modelo)
        if not os.path.exists(caminho_modelo):
            print(f'   AVISO: modelo {modelo} nao encontrado em DOCS_MODELOS/')
            continue
        saida = os.path.join(base, CONTRATACAO, f'{nome.title()} - {titulo_doc}.docx')
        faltando = preencher_modelo(caminho_modelo, campos, saida)
        gerados.append({'documento': titulo_doc, 'arquivo': saida, 'campos_faltando': sorted(faltando),
                        'pendencias': pendencias_no_arquivo(saida)})
    return gerados, campos


def pendencias_no_arquivo(caminho):
    """Tudo que ainda esta marcado [PREENCHER ...] ou [CONFERIR ...] no documento."""
    doc = Document(caminho)
    textos = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for linha in t.rows:
            textos.extend(c.text for c in linha.cells)
    achados = []
    for t in textos:
        achados.extend(PENDENCIA.findall(t))
    return sorted(set(achados))
