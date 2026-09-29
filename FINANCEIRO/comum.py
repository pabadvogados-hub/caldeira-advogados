"""
Funcoes comuns do FINANCEIRO: formatacao, datas, lista de clientes que nao recebem
cobranca, planilhas XLSX e consultas ao Asaas que nao existem em INTEGRACOES/.

Consultas extras ao Asaas (so leitura), usando o _request de INTEGRACOES/asaas_integration.py:
- cobrancas_recebidas(inicio, fim): cobrancas com DATA DE PAGAMENTO no periodo
- obter_cobranca(id)
- cobrancas_do_cpf(cpf): todas as cobrancas dos clientes Asaas com aquele CPF
"""
import os
import re
import sys
import unicodedata

sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'NUCLEO'))
import ambiente  # noqa: E402  (carrega .env e caminhos)
from datetime import date, datetime  # noqa: E402

from config.escritorio import ESCRITORIO, VISUAL  # noqa: E402

AQUI = os.path.dirname(os.path.abspath(__file__))
ARQ_NAO_COBRAR = os.path.join(AQUI, 'clientes_nao_cobrar.txt')
NOME_ESCRITORIO = ESCRITORIO['nome']

MESES = {1: 'JANEIRO', 2: 'FEVEREIRO', 3: 'MARCO', 4: 'ABRIL', 5: 'MAIO', 6: 'JUNHO', 7: 'JULHO',
         8: 'AGOSTO', 9: 'SETEMBRO', 10: 'OUTUBRO', 11: 'NOVEMBRO', 12: 'DEZEMBRO'}

# status do Asaas que significam dinheiro recebido
STATUS_RECEBIDO = ('RECEIVED', 'CONFIRMED', 'RECEIVED_IN_CASH')


# ============================================================
# FORMATACAO E DATAS
# ============================================================

def moeda(v):
    return f'R$ {num(v):,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.')


def pct(v):
    return f'{v:.1f}%'.replace('.', ',')


def num(v):
    """Valor em float, aceitando 1234.5 / '1234.50' / '1.234,50' / 'R$ 1.234,50'."""
    if v in (None, ''):
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace('R$', '').replace(' ', '')
    if ',' in s:
        s = s.replace('.', '').replace(',', '.')
    try:
        return float(s)
    except ValueError:
        return 0.0


def para_data(v):
    """date a partir de 'AAAA-MM-DD', 'DD/MM/AAAA' ou date. None se nao der."""
    if isinstance(v, date):
        return v
    for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%Y-%m-%d %H:%M:%S'):
        try:
            return datetime.strptime(str(v)[:19], fmt).date()
        except (TypeError, ValueError):
            continue
    return None


def data_br(v):
    d = para_data(v)
    return d.strftime('%d/%m/%Y') if d else (v or '-')


def dias_do_vencimento(vencimento, hoje=None):
    """Dias em relacao ao vencimento: negativo = faltam, 0 = hoje, positivo = atraso."""
    d = para_data(vencimento)
    if not d:
        return None
    return ((hoje or date.today()) - d).days


def sem_acento(s):
    return ''.join(c for c in unicodedata.normalize('NFD', s or '') if unicodedata.category(c) != 'Mn')


def digitos(s):
    return ''.join(c for c in str(s or '') if c.isdigit())


def normalizar_nome(nome):
    return re.sub(r'\s+', ' ', sem_acento(str(nome or '')).upper()).strip()


def nomes_parecidos(a, b):
    """Mesmo cliente? Igual, um contido no outro ou mesmo primeiro + ultimo nome."""
    a, b = normalizar_nome(a), normalizar_nome(b)
    if not a or not b:
        return False
    if a == b or a in b or b in a:
        return True
    pa, pb = a.split(), b.split()
    return len(pa) >= 2 and len(pb) >= 2 and pa[0] == pb[0] and pa[-1] == pb[-1]


def primeiro_nome(nome):
    partes = (nome or '').split()
    return partes[0].title() if partes else 'tudo bem'


def pasta_saida(saida=None, *partes):
    """Pasta de saida: --saida (testes) ou SAIDA/FINANCEIRO/... (padrao, fora do git)."""
    if saida:
        caminho = os.path.join(saida, *partes)
        os.makedirs(caminho, exist_ok=True)
        return caminho
    return ambiente.pasta_saida('FINANCEIRO', *partes)


# ============================================================
# CLIENTES QUE NAO RECEBEM COBRANCA AUTOMATICA
# ============================================================

def carregar_nao_cobrar(caminho=ARQ_NAO_COBRAR):
    lista = {'ids': set(), 'cpfs': set(), 'nomes': []}
    if not os.path.isfile(caminho):
        return lista
    with open(caminho, encoding='utf-8') as f:
        for linha in f:
            linha = linha.strip()
            if not linha or linha.startswith('#'):
                continue
            if linha.startswith('cus_'):
                lista['ids'].add(linha)
            elif len(digitos(linha)) in (11, 14) and not re.search(r'[A-Za-z]', linha):
                lista['cpfs'].add(digitos(linha))
            else:
                lista['nomes'].append(normalizar_nome(linha))
    return lista


def esta_na_lista(lista, customer_id=None, nome=None, cpf=None):
    if customer_id and customer_id in lista['ids']:
        return True
    if cpf and digitos(cpf) in lista['cpfs']:
        return True
    n = normalizar_nome(nome)
    return bool(n) and any(x and x in n for x in lista['nomes'])


# ============================================================
# ASAAS (so leitura)
# ============================================================

def asaas_configurado():
    return ambiente.tem_credencial('ASAAS_API_TOKEN')


def advbox_configurado():
    return ambiente.tem_credencial('ADVBOX_API_TOKEN')


class ClientesAsaas:
    """Cache de clientes do Asaas (nome, CPF, telefone) para nao consultar o mesmo duas vezes."""

    def __init__(self, fixos=None):
        self._cache = dict(fixos or {})

    def get(self, customer_id):
        if customer_id not in self._cache:
            try:
                from asaas_integration import obter_cliente
                self._cache[customer_id] = obter_cliente(customer_id) or {}
            except Exception as e:  # noqa: BLE001
                print(f'   AVISO: cliente {customer_id} nao consultado no Asaas ({e})')
                self._cache[customer_id] = {}
        return self._cache[customer_id]

    @staticmethod
    def telefone(cli):
        return cli.get('mobilePhone') or cli.get('phone') or ''


def _paginar(endpoint, params):
    from asaas_integration import _request
    todos, offset = [], 0
    while True:
        p = dict(params, limit=100, offset=offset)
        data = _request('GET', endpoint, params=p) or {}
        registros = data.get('data', [])
        todos.extend(registros)
        if not data.get('hasMore') or not registros:
            break
        offset += 100
    return todos


def cobrancas_recebidas(inicio, fim):
    """Cobrancas com pagamento entre inicio e fim (AAAA-MM-DD), so status de recebido."""
    pagos = _paginar('/payments', {'paymentDate[ge]': inicio, 'paymentDate[le]': fim})
    return [p for p in pagos if p.get('status') in STATUS_RECEBIDO and not p.get('deleted')]


def cobrancas_abertas(status, venc_de, venc_ate):
    from asaas_integration import listar_cobrancas
    return [p for p in listar_cobrancas(status=status, due_date_ge=venc_de, due_date_le=venc_ate)
            if not p.get('deleted')]


def obter_cobranca(payment_id):
    from asaas_integration import _request
    try:
        return _request('GET', f'/payments/{payment_id}') or {}
    except Exception:  # noqa: BLE001
        return {}


def cobrancas_do_cpf(cpf):
    from asaas_integration import listar_clientes
    cpf = digitos(cpf)
    if not cpf:
        return []
    cobr = []
    for cli in listar_clientes(cpf=cpf):
        cobr += [p for p in _paginar('/payments', {'customer': cli['id']}) if not p.get('deleted')]
    return cobr


# ============================================================
# PLANILHAS (XLSX)
# ============================================================

def _openpyxl():
    try:
        import openpyxl
        return openpyxl
    except ImportError:
        raise SystemExit('ERRO: falta o pacote openpyxl. Rode: pip install openpyxl '
                         '(ou deploy\\instalar_windows.bat).')


def planilha_nova():
    wb = _openpyxl().Workbook()
    wb.remove(wb.active)
    return wb


def aba(wb, titulo, cabecalho, linhas, larguras=None, colunas_moeda=(), titulo_topo=None):
    """Cria uma aba com cabecalho na cor do escritorio. colunas_moeda: indices (0-based)."""
    op = _openpyxl()
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    ws = wb.create_sheet(titulo[:31])
    borda = Border(*(Side('thin', color='BFBFBF'),) * 4)
    linha0 = 1
    if titulo_topo:
        ws.cell(row=1, column=1, value=titulo_topo).font = Font(bold=True, size=13, color=VISUAL['cor_destaque'])
        linha0 = 3
    for i, nome in enumerate(cabecalho, 1):
        c = ws.cell(row=linha0, column=i, value=nome)
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', start_color=VISUAL['cor_destaque'], end_color=VISUAL['cor_destaque'])
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = borda
    for r, linha in enumerate(linhas, linha0 + 1):
        for i, valor in enumerate(linha, 1):
            c = ws.cell(row=r, column=i, value=valor)
            c.border = borda
            if (i - 1) in colunas_moeda and isinstance(valor, (int, float)):
                c.number_format = '#,##0.00'
    for i, w in enumerate(larguras or [], 1):
        ws.column_dimensions[op.utils.get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=linha0 + 1, column=1)
    return ws


def salvar_planilha(wb, caminho):
    wb.save(caminho)
    return caminho


def cabecalho_execucao(titulo, exemplo=False):
    print('=' * 72)
    print(f'{titulo} - {NOME_ESCRITORIO}')
    print(f"Executado em {datetime.now().strftime('%d/%m/%Y %H:%M')}"
          + ('   [EXEMPLO: dados ficticios, nada e consultado nem enviado]' if exemplo else ''))
    print('=' * 72)
