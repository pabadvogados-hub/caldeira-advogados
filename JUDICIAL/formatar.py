"""
Converte o texto com marcacao simples (o que a IA devolve) em .docx no timbrado do escritorio.

Marcacao (uma instrucao por linha):
    !! TEXTO            enderecamento (linhas seguidas viram um paragrafo so)
    == TEXTO            titulo da peca / nome da acao, centralizado (linhas seguidas se juntam)
    # TEXTO             titulo de secao (I. DA GRATUIDADE ...)
    ## TEXTO            subtitulo (1) DAS CIRCUNSTANCIAS ...)
    ### TEXTO           sub-subtitulo (a) Estiagem severa ...)
    > TEXTO             citacao: recuo 4 cm, italico, fonte menor, espacamento simples
    - TEXTO             item com marcador
    | a | b |           tabela (linha |---| e ignorada; a primeira linha e o cabecalho)
    @@ TEXTO            legenda centralizada (IMAGEM 01 - ...)
    :: TEXTO            paragrafo centralizado (fecho, local e data)
    @@REVISAO           aviso "PRONTA PARA REVISAO DO(A) ADVOGADO(A) RESPONSAVEL"
    @@ASSINATURAS       bloco de assinaturas dos advogados (config/escritorio.py OUTORGADOS)
    qualquer outra      paragrafo do corpo: justificado, 1,5, recuo de primeira linha

No meio do texto: **negrito**, ^^destaque^^ (negrito + laranja, para nomes das partes e valores),
*italico*. Marcas [CONFERIR ...] e [PREENCHER ...] saem sempre em vermelho e negrito.

Perfis (tirados das pecas reais, DNA secao 8.1):
    'inicial'  enderecamento Times 14 laranja; titulos de secao Times 14 negrito laranja a direita;
               subtitulos laranja; nomes das partes e valores em laranja.
    'sobrio'   replica, embargos, contrarrazoes, manifestacao: tudo preto, titulos negrito 12.
"""
import os
import re
import sys
from datetime import date

sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'NUCLEO'))
import ambiente  # noqa: E402,F401  (carrega .env e caminhos)

from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Pt, RGBColor  # noqa: E402

from docx_caldeira import DESTAQUE, TEXTO, novo_documento, tabela  # noqa: E402
from config.escritorio import OUTORGADOS, VISUAL  # noqa: E402

VERMELHO = RGBColor(0xC0, 0x00, 0x00)
PRETO = RGBColor(0, 0, 0)
MARCA = re.compile(r'(\[(?:CONFERIR|PREENCHER)[^\]]*\])')
PENDENCIA = re.compile(r'\[(?:CONFERIR|PREENCHER)[^\]]*\]')
INLINE = re.compile(r'(\*\*.+?\*\*|\^\^.+?\^\^|(?<![\w*])\*[^*\s][^*]*?\*(?![\w*]))')

AVISO_REVISAO = 'PRONTA PARA REVISÃO DO(A) ADVOGADO(A) RESPONSÁVEL'

PERFIS = {
    'inicial': {'cor': True, 'tam_enderecamento': 14, 'tam_secao': 14, 'alinha_secao': 'direita'},
    'sobrio': {'cor': False, 'tam_enderecamento': 12, 'tam_secao': 12, 'alinha_secao': 'esquerda'},
}


# ============================================================
# PARAGRAFOS E TRECHOS
# ============================================================

def _fmt(p, alinhamento=WD_ALIGN_PARAGRAPH.JUSTIFY, recuo=False, espaco=None, antes=0, depois=6):
    p.alignment = alinhamento
    f = p.paragraph_format
    f.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    f.line_spacing = espaco or VISUAL['espacamento']
    f.space_before = Pt(antes)
    f.space_after = Pt(depois)
    f.first_line_indent = Cm(VISUAL['recuo_primeira_linha_cm']) if recuo else Cm(0)
    return p


def _run(p, texto, negrito=False, italico=False, cor=None, tamanho=None):
    """Escreve um trecho, destacando em vermelho as marcas [CONFERIR]/[PREENCHER]."""
    for parte in MARCA.split(texto):
        if not parte:
            continue
        r = p.add_run(parte)
        r.bold, r.italic = negrito, italico
        if tamanho:
            r.font.size = Pt(tamanho)
        if MARCA.fullmatch(parte):
            r.bold, r.italic, r.font.color.rgb = True, False, VERMELHO
        elif cor is not None:
            r.font.color.rgb = cor


def _inline(p, texto, perfil, negrito=False, italico=False, cor=None, tamanho=None):
    """Interpreta **negrito**, ^^destaque^^ e *italico* dentro de uma linha (aceita um dentro do outro)."""
    for parte in INLINE.split(texto):
        if not parte:
            continue
        if parte.startswith('**') and parte.endswith('**') and len(parte) > 4:
            _inline(p, parte[2:-2], perfil, True, italico, cor, tamanho)
        elif parte.startswith('^^') and parte.endswith('^^') and len(parte) > 4:
            _inline(p, parte[2:-2], perfil, True, italico, DESTAQUE if perfil['cor'] else cor, tamanho)
        elif parte.startswith('*') and parte.endswith('*') and len(parte) > 2:
            _inline(p, parte[1:-1], perfil, negrito, True, cor, tamanho)
        else:
            _run(p, parte, negrito, italico, cor, tamanho)


def _limpa_inline(texto):
    return texto.replace('**', '').replace('^^', '')


def _maiusculo(texto):
    """Caixa alta, menos dentro das marcas [CONFERIR ...]/[PREENCHER ...]."""
    return ''.join(parte if MARCA.fullmatch(parte) else parte.upper() for parte in MARCA.split(texto))


def _sem_bordas(t):
    tblPr = t._tbl.tblPr
    bordas = OxmlElement('w:tblBorders')
    for lado in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement(f'w:{lado}')
        el.set(qn('w:val'), 'nil')
        bordas.append(el)
    tblPr.append(bordas)


def assinaturas(doc, perfil):
    cor = DESTAQUE if perfil['cor'] else PRETO
    advs = OUTORGADOS or []
    if len(advs) <= 1:
        for a in advs:
            for i, linha in enumerate((a['nome'], a['oab'])):
                p = _fmt(doc.add_paragraph(), WD_ALIGN_PARAGRAPH.CENTER, espaco=1.0, antes=24 if i == 0 else 0, depois=0)
                _run(p, linha, negrito=(i == 0), cor=cor if i == 0 else None)
        return
    doc.add_paragraph()
    t = doc.add_table(rows=1, cols=len(advs))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    _sem_bordas(t)
    for i, a in enumerate(advs):
        cel = t.rows[0].cells[i]
        cel.text = ''
        p1 = _fmt(cel.paragraphs[0], WD_ALIGN_PARAGRAPH.CENTER, espaco=1.0, antes=18, depois=0)
        _run(p1, a['nome'], negrito=True, cor=cor)
        p2 = _fmt(cel.add_paragraph(), WD_ALIGN_PARAGRAPH.CENTER, espaco=1.0, depois=0)
        _run(p2, a['oab'])


def aviso_revisao(doc):
    p = _fmt(doc.add_paragraph(), WD_ALIGN_PARAGRAPH.CENTER, espaco=1.15, antes=12, depois=12)
    _run(p, AVISO_REVISAO, negrito=True, cor=VERMELHO)
    p.runs[-1].font.color.rgb = VERMELHO
    p2 = _fmt(doc.add_paragraph(), WD_ALIGN_PARAGRAPH.CENTER, espaco=1.0, depois=12)
    r = p2.add_run(f'Minuta gerada pela IA em {date.today().strftime("%d/%m/%Y")}. Apagar este aviso e todas as '
                   'marcas em vermelho antes de protocolar. A IA não protocola.')
    r.italic, r.font.size, r.font.color.rgb = True, Pt(9), VERMELHO


# ============================================================
# CONVERSAO
# ============================================================

def _tabela_md(doc, linhas, perfil):
    celulas = []
    for l in linhas:
        partes = [c.strip() for c in l.strip().strip('|').split('|')]
        if all(re.fullmatch(r':?-{2,}:?', c or '--') for c in partes):
            continue  # linha separadora |---|---|
        celulas.append([_limpa_inline(c) for c in partes])
    if not celulas:
        return
    n = max(len(c) for c in celulas)
    celulas = [c + [''] * (n - len(c)) for c in celulas]
    tabela(doc, celulas[0], celulas[1:], tamanho=9)


def _bloco(doc, tipo, texto, perfil):
    cor_titulo = DESTAQUE if perfil['cor'] else PRETO
    if tipo == 'enderecamento':
        p = _fmt(doc.add_paragraph(), WD_ALIGN_PARAGRAPH.JUSTIFY, depois=24)
        _inline(p, _maiusculo(texto), perfil, negrito=True, cor=cor_titulo, tamanho=perfil['tam_enderecamento'])
    elif tipo == 'titulo':
        p = _fmt(doc.add_paragraph(), WD_ALIGN_PARAGRAPH.CENTER, antes=12, depois=12)
        _inline(p, _maiusculo(texto), perfil, negrito=True, cor=cor_titulo)
    elif tipo == 'secao':
        alinha = WD_ALIGN_PARAGRAPH.RIGHT if perfil['alinha_secao'] == 'direita' else WD_ALIGN_PARAGRAPH.LEFT
        p = _fmt(doc.add_paragraph(), alinha, antes=14, depois=8)
        p.paragraph_format.keep_with_next = True
        _inline(p, _maiusculo(_limpa_inline(texto)), perfil, negrito=True, cor=cor_titulo, tamanho=perfil['tam_secao'])
    elif tipo == 'subtitulo':
        p = _fmt(doc.add_paragraph(), WD_ALIGN_PARAGRAPH.LEFT, antes=10, depois=6)
        p.paragraph_format.keep_with_next = True
        _inline(p, _limpa_inline(texto), perfil, negrito=True, cor=cor_titulo)
    elif tipo == 'subsub':
        p = _fmt(doc.add_paragraph(), WD_ALIGN_PARAGRAPH.LEFT, antes=6, depois=4)
        p.paragraph_format.keep_with_next = True
        _inline(p, _limpa_inline(texto), perfil, negrito=True, cor=TEXTO if perfil['cor'] else PRETO)
    elif tipo == 'citacao':
        p = _fmt(doc.add_paragraph(), WD_ALIGN_PARAGRAPH.JUSTIFY, espaco=1.0, antes=4, depois=10)
        p.paragraph_format.left_indent = Cm(4)
        _inline(p, texto, perfil, italico=True, tamanho=10)
    elif tipo == 'item':
        p = _fmt(doc.add_paragraph(), depois=3)
        p.paragraph_format.left_indent = Cm(1.25)
        p.paragraph_format.first_line_indent = Cm(-0.5)
        _inline(p, '•  ' + texto, perfil)
    elif tipo == 'legenda':
        p = _fmt(doc.add_paragraph(), WD_ALIGN_PARAGRAPH.CENTER, espaco=1.0, antes=6, depois=6)
        _inline(p, texto, perfil, negrito=True, cor=cor_titulo, tamanho=11)
    elif tipo == 'centro':
        p = _fmt(doc.add_paragraph(), WD_ALIGN_PARAGRAPH.CENTER, depois=6)
        _inline(p, texto, perfil)
    elif ROTULO_SEM_RECUO.match(texto):
        p = _fmt(doc.add_paragraph(), WD_ALIGN_PARAGRAPH.LEFT, depois=4)
        _inline(p, texto, perfil)
    else:
        p = _fmt(doc.add_paragraph(), recuo=True)
        _inline(p, texto, perfil)


# linhas de identificacao (sem recuo): "**PROCESSO Nº ...**", "**APELANTE: ...**", "EGRÉGIO TRIBUNAL ..."
ROTULO_SEM_RECUO = re.compile(r'^(\*\*|\^\^)?\s*(PROCESSO|AUTOS|APELANTE|APELAD[OA]|AGRAVANTE|AGRAVAD[OA]|'
                              r'EMBARGANTE|EMBARGAD[OA]|DISTRIBUI[ÇC][ÃA]O POR DEPEND|EGR[ÉE]GIO|COLENDA|[ÍI]NCLITOS)\b')


PREFIXOS = (('!!', 'enderecamento'), ('###', 'subsub'), ('##', 'subtitulo'), ('#', 'secao'),
            ('==', 'titulo'), ('>', 'citacao'), ('::', 'centro'))
JUNTAR = ('enderecamento', 'titulo', 'citacao')  # linhas seguidas do mesmo tipo viram um paragrafo


def _classificar(linha):
    s = linha.strip()
    if s.startswith('%%') or (s.startswith('{{') and s.endswith('}}')):
        return 'ignorar', ''   # metadado do arquivo ou instrucao de esqueleto que a IA copiou
    if s in ('@@REVISAO', '@@ASSINATURAS'):
        return s[2:].lower(), ''
    if s.startswith('@@'):
        return 'legenda', s[2:].strip()
    if s.startswith('|') and s.endswith('|'):
        return 'tabela', s
    for pre, tipo in PREFIXOS:
        if s.startswith(pre):
            return tipo, s[len(pre):].strip().rstrip('=').strip() if tipo == 'titulo' else s[len(pre):].strip()
    if re.match(r'^[-•]\s+', s):
        return 'item', re.sub(r'^[-•]\s+', '', s)
    return 'corpo', s


def montar_docx(texto, saida, perfil='inicial'):
    """Gera o .docx a partir do texto marcado. Retorna a lista de pendencias [CONFERIR]/[PREENCHER]."""
    perfil_cfg = PERFIS.get(perfil, PERFIS['sobrio'])
    doc = novo_documento()
    atual, acumulado, tabela_linhas = None, [], []

    def descarregar():
        nonlocal atual, acumulado, tabela_linhas
        if tabela_linhas:
            _tabela_md(doc, tabela_linhas, perfil_cfg)
            tabela_linhas = []
        if atual and acumulado:
            _bloco(doc, atual, ' '.join(acumulado), perfil_cfg)
        atual, acumulado = None, []

    for linha in texto.replace('\r\n', '\n').split('\n'):
        if not linha.strip():
            descarregar()
            continue
        tipo, conteudo = _classificar(linha)
        if tipo == 'ignorar':
            continue
        if tipo == 'tabela':
            if atual:
                _bloco(doc, atual, ' '.join(acumulado), perfil_cfg)
                atual, acumulado = None, []
            tabela_linhas.append(conteudo)
            continue
        if tabela_linhas:
            descarregar()
        if tipo == 'citacao' and not conteudo:
            descarregar()  # linha "> " vazia separa paragrafos da citacao
            continue
        if tipo in JUNTAR and tipo == atual:
            acumulado.append(conteudo)
            continue
        descarregar()
        if tipo == 'revisao':
            aviso_revisao(doc)
        elif tipo == 'assinaturas':
            assinaturas(doc, perfil_cfg)
        elif tipo in JUNTAR:
            atual, acumulado = tipo, [conteudo]
        else:
            _bloco(doc, tipo, conteudo, perfil_cfg)
    descarregar()
    os.makedirs(os.path.dirname(os.path.abspath(saida)), exist_ok=True)
    doc.save(saida)
    return pendencias(texto)


def pendencias(texto):
    """Marcas [CONFERIR ...] / [PREENCHER ...] que continuam no texto."""
    return PENDENCIA.findall(texto or '')


def texto_plano(texto):
    """Texto sem a marcacao (para mostrar trechos no terminal)."""
    saida = []
    for linha in texto.split('\n'):
        tipo, conteudo = _classificar(linha) if linha.strip() else ('corpo', '')
        saida.append(_limpa_inline(conteudo if tipo != 'tabela' else linha))
    return '\n'.join(saida)
