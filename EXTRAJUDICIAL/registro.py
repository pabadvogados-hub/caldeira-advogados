"""
Estado da FASE EXTRAJUDICIAL no caso.json do cliente (chave 'extrajudicial').

caso['extrajudicial'] = {
  'notificacoes': [ {                       # uma por notificacao (a reiteracao ao mesmo banco e outra)
      'banco', 'tipo' ('alongamento' | 'contratos'), 'numero' (1a, 2a... ao mesmo banco),
      'arquivo' (.docx), 'pdf', 'gerada_em', 'pendencias',
      'email_destino', 'rascunho_id', 'rascunho_em', 'anexos',
      'enviada_em', 'prazo_resposta',
      'resposta': {'tipo': 'recebida' | 'sem_resposta', 'data', 'arquivo', 'trecho', 'registrada_em'} | None,
      'propostas': [{'arquivo', 'parecer', 'recebida_em', 'gerado_em'}],
      'consumidor_gov': {'arquivo', 'gerado_em'} | None,
      'decisao': {'resultado' ('acordo' | 'sem-acordo' | 'judicial'), 'data', 'observacao'} | None,
      'whatsapp_cliente_em', 'proxima_acao', 'historico': [{'data', 'evento'}],
  } ],
  'relatorios': [{'arquivo', 'gerado_em'}],
  'ultimo_acompanhamento': 'AAAA-MM-DD',
}
Datas gravadas em ISO (AAAA-MM-DD); na tela e nos documentos, DD/MM/AAAA.
"""
import os
import re
import sys
import unicodedata

sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'NUCLEO'))
import ambiente  # noqa: E402,F401  (carrega .env e caminhos)
from datetime import date, datetime, timedelta  # noqa: E402

from docx import Document  # noqa: E402

from CONTRATACAO import pasta_cliente  # noqa: E402
from CONTRATACAO.extrator import ler_arquivo  # noqa: E402
from config.escritorio import CHECKLIST_DOCUMENTOS, PASTA_AREA, SUBPASTAS_CLIENTE  # noqa: E402

# Prazo que o escritorio da ao banco para responder (dias corridos a partir do envio).
# A notificacao nao fixa prazo no texto; este prazo e so do controle interno.
PRAZO_RESPOSTA_BANCO_DIAS = int(os.getenv('PRAZO_RESPOSTA_BANCO_DIAS') or 10)
ALERTA_INICIAL_DIAS = 15      # avisa quando faltarem ate 15 dias para o limite da inicial
ALERTA_PARADO_DIAS = 2        # minuta/rascunho parado sem envio

PENDENCIA = re.compile(r'\[(?:PREENCHER|CONFERIR)[^\]]*\]')
PASTA_EXTRA = next((s for s in SUBPASTAS_CLIENTE if 'EXTRAJUDICIAL' in s.upper()), '10 EXTRAJUDICIAL')
MESES = ['janeiro', 'fevereiro', 'março', 'abril', 'maio', 'junho', 'julho', 'agosto', 'setembro',
         'outubro', 'novembro', 'dezembro']
TIPOS = {'alongamento': 'PEDIDO DE ALONGAMENTO DE DÍVIDA RURAL', 'contratos': 'PEDIDO DE CONTRATOS RURAIS'}


# ============================================================
# DATAS
# ============================================================

def hoje():
    """Data de hoje. HOJE_SIMULADO=AAAA-MM-DD no ambiente serve so para teste (simular prazos)."""
    simulado = os.getenv('HOJE_SIMULADO')
    return date.fromisoformat(simulado) if simulado else date.today()


def iso(d):
    return d.isoformat() if d else None


def de_iso(txt):
    try:
        return date.fromisoformat((txt or '')[:10])
    except ValueError:
        return None


def de_br(txt):
    try:
        return datetime.strptime((txt or '').strip()[:10], '%d/%m/%Y').date()
    except ValueError:
        return None


def br(d_ou_iso):
    d = de_iso(d_ou_iso) if isinstance(d_ou_iso, str) else d_ou_iso
    return d.strftime('%d/%m/%Y') if d else '-'


def extenso(d=None):
    d = d or hoje()
    return f'{d.day} de {MESES[d.month - 1]} de {d.year}'


# ============================================================
# NOMES DE BANCO
# ============================================================

def norm(txt):
    s = ''.join(c for c in unicodedata.normalize('NFD', txt or '') if unicodedata.category(c) != 'Mn')
    return re.sub(r'[^A-Z0-9 ]+', ' ', s.upper()).split()


def mesmo_banco(a, b):
    """'Sicredi' casa com 'SICREDI UNIVALES'; 'BB' so casa com 'BB'."""
    na, nb = ' '.join(norm(a)), ' '.join(norm(b))
    if not na or not nb:
        return False
    return na == nb or (len(na) > 3 and na in nb) or (len(nb) > 3 and nb in na)


def arquivo_seguro(txt):
    limpo = re.sub(r'[<>:"/\\|?*]', '', txt or '').strip()
    return re.sub(r'\s+', ' ', limpo)


# ============================================================
# CASO
# ============================================================

def abrir_caso(pasta):
    """Aceita o caminho da pasta do cliente ou so o nome (procura em PASTA_CLIENTES_RAIZ/AGRONEGOCIO)."""
    base = pasta
    if not os.path.isdir(base):
        base = os.path.join(pasta_cliente.raiz_clientes(), PASTA_AREA, pasta_cliente.nome_pasta(pasta))
    caso = pasta_cliente.ler_caso(base)
    if not caso:
        raise SystemExit(f'ERRO: nao achei 00 CONTRATACAO/caso.json em {pasta}')
    return os.path.abspath(base), caso


def salvar(base, caso):
    pasta_cliente.salvar_caso(base, caso)


def extrajudicial(caso):
    ext = caso.setdefault('extrajudicial', {})
    ext.setdefault('notificacoes', [])
    ext.setdefault('relatorios', [])
    return ext


def nome_cliente(caso):
    return (caso.get('qualificacao') or {}).get('nome') or '[CONFERIR nome do cliente]'


def nome_arquivo_cliente(caso):
    return pasta_cliente.nome_pasta(nome_cliente(caso)).title()


def pasta_extra(base):
    caminho = os.path.join(base, PASTA_EXTRA)
    os.makedirs(caminho, exist_ok=True)
    return caminho


def bancos_do_caso(caso):
    vistos = []
    for o in (caso.get('triagem') or {}).get('operacoes') or []:
        b = (o.get('banco') or '').strip()
        if b and not any(mesmo_banco(b, v) for v in vistos):
            vistos.append(b)
    for n in extrajudicial(caso)['notificacoes']:
        if not any(mesmo_banco(n['banco'], v) for v in vistos):
            vistos.append(n['banco'])
    return vistos


def operacoes_do_banco(caso, banco):
    return [o for o in (caso.get('triagem') or {}).get('operacoes') or [] if mesmo_banco(o.get('banco'), banco)]


NAO_BANCARIO = re.compile(r'n[aã]o banc|cerealista|trading|fornecedor|revenda|armaz[eé]m', re.I)


def credor_nao_bancario(caso, banco):
    """Cerealista, trading, revenda de insumos: fora do credito rural bancario (so notifica com --banco)."""
    return bool(NAO_BANCARIO.search(banco or '') or
                any(NAO_BANCARIO.search(o.get('observacao') or '') for o in operacoes_do_banco(caso, banco)))


def bancos_ativos(caso):
    """Credores da fase extrajudicial: bancos/cooperativas, mais o nao bancario que ja foi notificado."""
    return [b for b in bancos_do_caso(caso) if notificacoes_do_banco(caso, b) or not credor_nao_bancario(caso, b)]


def resolver_banco(caso, texto):
    """Casa o --banco digitado com um credor do caso (tambem pelos apelidos da base de e-mails)."""
    achados = [b for b in bancos_do_caso(caso) if mesmo_banco(texto, b)]
    if not achados:
        import bancos as base_bancos   # aqui para nao criar import circular
        for entrada in base_bancos.entradas_do_banco(texto):
            achados = [b for b in bancos_do_caso(caso) if base_bancos.casa(entrada, b)]
            if achados:
                break
    if len(achados) == 1:
        return achados[0]
    lista = ', '.join(bancos_do_caso(caso)) or 'nenhum'
    if not achados:
        raise SystemExit(f'ERRO: "{texto}" nao e credor deste caso. Credores: {lista}')
    raise SystemExit(f'ERRO: "{texto}" casa com mais de um credor ({", ".join(achados)}). Seja mais especifico.')


def notificacoes_do_banco(caso, banco):
    lista = [n for n in extrajudicial(caso)['notificacoes'] if mesmo_banco(n['banco'], banco)]
    for n in lista:
        # registro antigo ou feito a mao: envio sem prazo calculado
        if n.get('enviada_em') and not n.get('prazo_resposta') and de_iso(n['enviada_em']):
            n['prazo_resposta'] = iso(de_iso(n['enviada_em']) + timedelta(days=PRAZO_RESPOSTA_BANCO_DIAS))
    return lista


def ultima(caso, banco):
    lista = notificacoes_do_banco(caso, banco)
    return lista[-1] if lista else None


def evento(n, texto):
    n.setdefault('historico', []).append({'data': iso(hoje()), 'evento': texto})


def nova_notificacao(caso, banco, tipo, arquivo, pdf, pendencias):
    anteriores = notificacoes_do_banco(caso, banco)
    n = {
        'banco': banco, 'tipo': tipo, 'numero': len(anteriores) + 1,
        'arquivo': arquivo, 'pdf': pdf, 'gerada_em': iso(hoje()), 'pendencias': pendencias,
        'email_destino': None, 'rascunho_id': None, 'rascunho_em': None, 'anexos': [],
        'enviada_em': None, 'prazo_resposta': None, 'resposta': None, 'propostas': [],
        'consumidor_gov': None, 'decisao': None, 'whatsapp_cliente_em': None,
        'proxima_acao': None, 'historico': [],
    }
    evento(n, f'Notificacao {n["numero"]}a ({tipo}) gerada')
    extrajudicial(caso)['notificacoes'].append(n)
    return n


def registrar_envio(n, data_envio):
    n['enviada_em'] = iso(data_envio)
    n['prazo_resposta'] = iso(data_envio + timedelta(days=PRAZO_RESPOSTA_BANCO_DIAS))
    evento(n, f'Enviada ao banco em {br(data_envio)} (confirmado pelo advogado)')


# ============================================================
# DOCUMENTOS
# ============================================================

def pendencias_docx(caminho):
    """Marcas [CONFERIR ...]/[PREENCHER ...] que ainda estao no .docx (texto e tabelas)."""
    if not caminho or not os.path.exists(caminho):
        return ['[CONFERIR arquivo nao encontrado]']
    doc = Document(caminho)
    textos = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for linha in t.rows:
            textos.extend(c.text for c in linha.cells)
    achados = []
    for t in textos:
        achados.extend(PENDENCIA.findall(t))
    return sorted(set(achados))


def pasta_item(base, item_id):
    item = next((i for i in CHECKLIST_DOCUMENTOS if i['id'] == item_id), None)
    if not item or not item.get('pasta'):
        return None
    if item['pasta'] == pasta_cliente.CONTRATACAO:
        return os.path.join(base, item['pasta'])
    return os.path.join(base, pasta_cliente.DOCS, item['pasta'])


def arquivos_item(base, item_id):
    pasta = pasta_item(base, item_id)
    if not pasta or not os.path.isdir(pasta):
        return []
    return [os.path.join(pasta, f) for f in sorted(os.listdir(pasta))
            if f.lower().endswith(pasta_cliente.EXTENSOES_DOC)]


def texto_documentos(base, item_id, banco=None, limite_arquivo=30000, limite_total=90000):
    """Texto dos arquivos de uma subpasta do checklist. Com banco, arquivos que citam o banco vem primeiro."""
    lidos = []
    for caminho in arquivos_item(base, item_id):
        txt = ler_arquivo(caminho) or ''
        if txt.strip():
            lidos.append((os.path.basename(caminho), txt[:limite_arquivo]))
    if banco:
        lidos.sort(key=lambda x: 0 if mesmo_banco(banco, x[0]) or ' '.join(norm(banco)) in ' '.join(norm(x[1][:5000])) else 1)
    saida, total = [], 0
    for nome, txt in lidos:
        if total + len(txt) > limite_total:
            saida.append(f'\n--- DOCUMENTO: {nome} (nao lido: limite de texto) ---')
            continue
        saida.append(f'\n--- DOCUMENTO: {nome} ---\n{txt}')
        total += len(txt)
    return ''.join(saida)


def procuracao_assinada(base):
    pasta = os.path.join(base, pasta_cliente.CONTRATACAO)
    if not os.path.isdir(pasta):
        return None
    for f in sorted(os.listdir(pasta)):
        if 'procura' in f.lower() and 'assinad' in f.lower() and f.lower().endswith('.pdf'):
            return os.path.join(pasta, f)
    return None


# ============================================================
# SITUACAO, PRAZOS E PROXIMA ACAO
# ============================================================

def prazos_caso(caso):
    return {p.get('id'): de_br(p.get('data')) for p in caso.get('prazos') or []}


def situacao(n):
    d = (n.get('decisao') or {}).get('resultado')
    if d == 'acordo':
        return 'ACORDO'
    if d == 'judicial':
        return 'ENCAMINHADO AO JUDICIAL'
    if d == 'sem-acordo':
        return 'SEM ACORDO'
    if n.get('propostas'):
        return 'PROPOSTA EM ANALISE'
    r = n.get('resposta') or {}
    if r.get('tipo') == 'recebida':
        return 'RESPOSTA RECEBIDA'
    if r.get('tipo') == 'sem_resposta':
        return 'SEM RESPOSTA'
    if n.get('enviada_em'):
        prazo = de_iso(n.get('prazo_resposta'))
        return 'PRAZO VENCIDO' if prazo and hoje() > prazo else 'AGUARDANDO RESPOSTA'
    if n.get('rascunho_id'):
        return 'RASCUNHO NO GMAIL'
    return 'MINUTA'


def proxima_acao_banco(caso, banco):
    notifs = notificacoes_do_banco(caso, banco)
    if not notifs:
        return f'Gerar a notificação ao {banco} (notificar).'
    n = notifs[-1]
    st = situacao(n)
    if st == 'MINUTA':
        pend = n.get('pendencias') or []
        if pend:
            return (f'Revisar a minuta e resolver {len(pend)} marca(s) CONFERIR/PREENCHER no .docx; '
                    'depois criar o rascunho no Gmail (rascunho) ou enviar e rodar registrar-envio.')
        return 'Minuta sem pendências: criar o rascunho no Gmail (rascunho) ou enviar ao banco e rodar registrar-envio.'
    if st == 'RASCUNHO NO GMAIL':
        return 'Revisar o rascunho no Gmail e clicar em Enviar; depois rodar registrar-envio.'
    if st == 'AGUARDANDO RESPOSTA':
        return f"Aguardar a resposta até {br(n.get('prazo_resposta'))} e cobrar o retorno do banco (gerente/agência)."
    if st == 'PRAZO VENCIDO':
        return (f"Prazo de resposta venceu em {br(n.get('prazo_resposta'))}: cobrar o banco e registrar a resposta "
                '(registrar-resposta --arquivo) ou a ausência (registrar-resposta --sem-resposta).')
    if st == 'RESPOSTA RECEBIDA':
        if n.get('tipo') == 'contratos':
            return ('Contratos recebidos: guardar as cédulas em 03 CEDULAS E CONTRATOS BANCARIOS e gerar a '
                    'notificação de alongamento (notificar --tipo alongamento).')
        return ('Resposta recebida: repassar ao Gestor Jurídico. Com proposta, rodar "proposta" (parecer); '
                'se o banco negou, registrar "decisao --resultado sem-acordo".')
    if st == 'PROPOSTA EM ANALISE':
        return ('Parecer da proposta pronto: o Gestor Jurídico decide (reunião com o gerente, se preciso) e '
                'registra "decisao --resultado acordo" ou "sem-acordo". Nada é aceito sem o Gestor.')
    if st == 'ACORDO':
        return 'Acordo aprovado pelo Gestor: formalizar o acordo extrajudicial, avisar o cliente e atualizar o ADVBOX.'
    if st == 'ENCAMINHADO AO JUDICIAL':
        return 'Com o Adv. Judicial: inicial com a notificação, o comprovante de envio e a resposta/ausência do banco.'
    # SEM RESPOSTA ou SEM ACORDO
    if n.get('tipo') == 'contratos' and st == 'SEM RESPOSTA':
        return ('Banco não enviou os contratos: notificação de alongamento (notificar --tipo alongamento, que reitera '
                'o pedido de cópias) + reclamação no consumidor.gov.br; na ação, pedir exibição (arts. 396 a 400 do CPC).')
    faltam = []
    if not any(x.get('consumidor_gov') for x in notifs):
        faltam.append('reclamação no consumidor.gov.br (consumidor-gov)')
    alongamentos = [x for x in notifs if x.get('tipo') == 'alongamento']
    if not alongamentos:
        faltam.append('notificação de alongamento (notificar --tipo alongamento)')
    elif len(alongamentos) < 2:
        faltam.append('nova notificação reiterando o pedido (notificar)')
    if faltam:
        return 'Sem acordo: ' + ' + '.join(faltam) + '; depois encaminhar ao judicial.'
    return ('Via extrajudicial esgotada: relatório ao Gestor (relatorio-gestor) e encaminhar ao Adv. Judicial '
            '(decisao --resultado judicial).')


def alertas_caso(caso):
    """Alertas de prazo do caso: notificacao atrasada, inicial perto do limite, resposta vencida, minuta parada."""
    alertas = []
    prazos = prazos_caso(caso)
    h = hoje()
    ext = extrajudicial(caso)
    bancos = bancos_ativos(caso)
    sem_envio =[b for b in bancos if not any(x.get('enviada_em') for x in notificacoes_do_banco(caso, b))]
    lim_notif = prazos.get('notificacao')
    if sem_envio and lim_notif:
        dias = (lim_notif - h).days
        if dias < 0:
            alertas.append(f'NOTIFICAÇÃO ATRASADA (limite {br(lim_notif)}): {", ".join(sem_envio)}')
        elif dias <= 3:
            alertas.append(f'notificação vence em {dias}d ({br(lim_notif)}): {", ".join(sem_envio)}')
    lim_ini = prazos.get('inicial')
    resolvidos = all(situacao(ultima(caso, b) or {}) in ('ACORDO', 'ENCAMINHADO AO JUDICIAL')
                     for b in bancos) if bancos else False
    if lim_ini and not resolvidos:
        dias = (lim_ini - h).days
        if dias < 0:
            alertas.append(f'INICIAL VENCIDA (limite {br(lim_ini)}, {-dias} dia(s) atrás)')
        elif dias <= ALERTA_INICIAL_DIAS:
            alertas.append(f'inicial em {dias}d (limite {br(lim_ini)}): decidir se segue ao judicial')
    for n in ext['notificacoes']:
        st = situacao(n)
        if st == 'PRAZO VENCIDO':
            alertas.append(f"{n['banco']}: prazo de resposta vencido em {br(n.get('prazo_resposta'))}")
        elif st in ('MINUTA', 'RASCUNHO NO GMAIL'):
            desde = de_iso(n.get('rascunho_em') or n.get('gerada_em'))
            if desde and (h - desde).days >= ALERTA_PARADO_DIAS:
                alertas.append(f"{n['banco']}: {st.lower()} parada há {(h - desde).days} dia(s) sem envio")
    return alertas
