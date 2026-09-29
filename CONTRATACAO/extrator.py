"""
Leitura de qualquer arquivo de entrada da fase de contratacao:
transcricao da reuniao (PDF, DOCX, TXT, SRT/VTT do Meet) e documentos do cliente
(CNH, RG, cedulas, extratos). PDF escaneado e imagem passam por OCR quando o
Tesseract estiver instalado.
"""
import os
import re
import shutil
import sys


def caminho_tesseract():
    """Tesseract no Windows, no Mac (Homebrew) ou no Linux. TESSERACT_CMD no .env tem prioridade.
    Precisa tambem do idioma portugues: Windows (marcar no instalador), Mac `brew install tesseract-lang`,
    Linux `apt install tesseract-ocr-por`."""
    achado = os.getenv('TESSERACT_CMD') or shutil.which('tesseract')
    if achado:
        return achado
    # o agendador (launchd no Mac, Agendador no Windows) nem sempre tem o PATH do terminal
    candidatos = [
        os.path.join(os.getenv('ProgramFiles', r'C:\Program Files'), 'Tesseract-OCR', 'tesseract.exe'),
        os.path.join(os.getenv('LOCALAPPDATA', ''), 'Programs', 'Tesseract-OCR', 'tesseract.exe'),
        '/opt/homebrew/bin/tesseract',   # Mac Apple Silicon (M1/M2/M3...)
        '/usr/local/bin/tesseract',      # Mac Intel
        '/usr/bin/tesseract',
    ]
    return next((c for c in candidatos if os.path.exists(c)), None)


_avisou_sem_tesseract = False


def _abrir_heic(caminho):
    """Foto do iPhone (.heic). Usa o pillow-heif se instalado; no Mac, converte com o 'sips' do sistema."""
    from PIL import Image
    try:
        from pillow_heif import register_heif_opener
        register_heif_opener()
        return Image.open(caminho)
    except ImportError:
        pass
    if sys.platform == 'darwin' and shutil.which('sips'):
        import subprocess
        import tempfile
        with tempfile.TemporaryDirectory() as tmp:
            jpg = os.path.join(tmp, 'foto.jpg')
            subprocess.run(['sips', '-s', 'format', 'jpeg', caminho, '--out', jpg], capture_output=True, timeout=60)
            if os.path.exists(jpg):
                img = Image.open(jpg)
                img.load()
                return img
    raise RuntimeError('foto .heic: no Windows, instale "pip install pillow-heif" ou salve a foto como JPG')


def _ocr_imagem(caminho_ou_pixmap):
    global _avisou_sem_tesseract
    try:
        import pytesseract
        from PIL import Image
    except ImportError:
        return ''
    tess = caminho_tesseract()
    if not tess:
        if not _avisou_sem_tesseract:
            print('   Aviso: sem Tesseract, imagem/PDF escaneado nao foi lido (OCR). Ver docs/COMECE_AQUI.md.')
            _avisou_sem_tesseract = True
        return ''
    pytesseract.pytesseract.tesseract_cmd = tess
    try:
        if isinstance(caminho_ou_pixmap, str):
            if caminho_ou_pixmap.lower().endswith(('.heic', '.heif')):
                img = _abrir_heic(caminho_ou_pixmap)
            else:
                img = Image.open(caminho_ou_pixmap)
        else:
            import io
            img = Image.open(io.BytesIO(caminho_ou_pixmap.tobytes('png')))
        return pytesseract.image_to_string(img, lang='por')
    except Exception as e:
        print(f'   Aviso: OCR falhou ({e})')
        return ''


def _ler_pdf(caminho, max_paginas=60):
    import fitz
    doc = fitz.open(caminho)
    partes = []
    for i, pagina in enumerate(doc):
        if i >= max_paginas:
            partes.append(f'\n[... {len(doc) - max_paginas} paginas nao lidas ...]')
            break
        texto = pagina.get_text()
        if len(texto.strip()) < 40:
            # pagina escaneada: tenta OCR
            texto = _ocr_imagem(pagina.get_pixmap(dpi=200)) or texto
        partes.append(texto)
    return '\n'.join(partes)


def _ler_legenda(caminho):
    """SRT/VTT (transcricao do Meet ou do Whisper): tira numeracao e marcacoes de tempo."""
    with open(caminho, encoding='utf-8', errors='ignore') as f:
        linhas = f.read().splitlines()
    saida = []
    for linha in linhas:
        l = linha.strip()
        if not l or l.isdigit() or '-->' in l or l.upper() == 'WEBVTT':
            continue
        if not saida or saida[-1] != l:
            saida.append(l)
    return ' '.join(saida)


def ler_arquivo(caminho):
    """Retorna o texto de um arquivo. Nunca levanta excecao: devolve '' se nao conseguir."""
    ext = os.path.splitext(caminho)[1].lower()
    try:
        if ext == '.pdf':
            return _ler_pdf(caminho)
        if ext in ('.txt', '.md'):
            with open(caminho, encoding='utf-8', errors='ignore') as f:
                return f.read()
        if ext in ('.srt', '.vtt'):
            return _ler_legenda(caminho)
        if ext == '.docx':
            from docx import Document
            d = Document(caminho)
            texto = [p.text for p in d.paragraphs]
            for t in d.tables:
                for linha in t.rows:
                    texto.append(' | '.join(c.text for c in linha.cells))
            return '\n'.join(texto)
        if ext in ('.jpg', '.jpeg', '.png', '.webp', '.bmp', '.tif', '.tiff', '.heic', '.heif'):
            return _ocr_imagem(caminho)
    except Exception as e:
        print(f'   Aviso: nao consegui ler {os.path.basename(caminho)} ({e})')
    return ''


def ler_cadastro(caminho):
    """
    Le o CADASTRO.txt preenchido pelo Closer/SDR (formato CHAVE: valor por linha).
    Chaves sao normalizadas para minusculas com _ (ex.: 'Estado civil' -> 'estado_civil').
    """
    dados = {}
    if not caminho or not os.path.exists(caminho):
        return dados
    with open(caminho, encoding='utf-8', errors='ignore') as f:
        for linha in f:
            if ':' not in linha or linha.strip().startswith('#'):
                continue
            chave, valor = linha.split(':', 1)
            chave = re.sub(r'[^a-z0-9]+', '_', _sem_acento(chave.strip().lower())).strip('_')
            valor = valor.strip()
            if valor:
                dados[chave] = valor
    return dados


def _sem_acento(s):
    import unicodedata
    return ''.join(c for c in unicodedata.normalize('NFD', s) if unicodedata.category(c) != 'Mn')
