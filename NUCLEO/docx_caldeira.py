"""
Documentos .docx no timbrado do Caldeira Advogados Associados.

- montar_timbrado(): gera config/timbrado_modelo.docx (logo no topo, marca d'agua,
  rodape com contatos) a partir das imagens tiradas das pecas do escritorio.
- novo_documento(): abre o timbrado para escrever um documento novo.
- preencher_modelo(): troca os {{CAMPOS}} de um modelo .docx (inclusive quando o Word
  quebrou o campo em varios pedacos de texto).
- docx_para_pdf(): converte para PDF (Word via docx2pdf ou LibreOffice; Windows, Mac e Linux).
"""
import copy
import os
import re
import shutil
import subprocess
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt, RGBColor

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, RAIZ)
from config.escritorio import ESCRITORIO, VISUAL  # noqa: E402

CONFIG = os.path.join(RAIZ, 'config')
TIMBRADO = os.path.join(CONFIG, 'timbrado_modelo.docx')
LOGO = os.path.join(CONFIG, 'logo_caldeira.png')
MARCA = os.path.join(CONFIG, 'marca_dagua_caldeira.png')

DESTAQUE = RGBColor.from_string(VISUAL['cor_destaque'])
TEXTO = RGBColor.from_string(VISUAL['cor_texto'])
CONFERIR = '[CONFERIR]'


# ============================================================
# TIMBRADO
# ============================================================

def _imagem_atras_do_texto(run, caminho, largura_cm, altura_cm):
    """Insere imagem centralizada na pagina, atras do texto (marca d'agua)."""
    run.add_picture(caminho, width=Cm(largura_cm), height=Cm(altura_cm))
    inline = run._r.xpath('.//wp:inline')[0]
    cx, cy = int(Cm(largura_cm)), int(Cm(altura_cm))
    anchor = parse_xml(
        f'<wp:anchor {nsdecls("wp", "a", "pic", "r")} distT="0" distB="0" distL="0" distR="0" '
        f'simplePos="0" relativeHeight="0" behindDoc="1" locked="1" layoutInCell="1" allowOverlap="1">'
        f'<wp:simplePos x="0" y="0"/>'
        f'<wp:positionH relativeFrom="page"><wp:align>center</wp:align></wp:positionH>'
        f'<wp:positionV relativeFrom="page"><wp:align>center</wp:align></wp:positionV>'
        f'<wp:extent cx="{cx}" cy="{cy}"/><wp:effectExtent l="0" t="0" r="0" b="0"/>'
        f'<wp:wrapNone/><wp:docPr id="900" name="MarcaDagua"/><wp:cNvGraphicFramePr/>'
        f'</wp:anchor>'
    )
    anchor.append(copy.deepcopy(inline.find(qn('a:graphic'))))
    inline.getparent().replace(inline, anchor)


def _borda_superior(paragrafo, cor='A6A6A6'):
    pPr = paragrafo._p.get_or_add_pPr()
    bdr = OxmlElement('w:pBdr')
    top = OxmlElement('w:top')
    for k, v in (('w:val', 'single'), ('w:sz', '6'), ('w:space', '4'), ('w:color', cor)):
        top.set(qn(k), v)
    bdr.append(top)
    pPr.append(bdr)


def montar_timbrado(saida=TIMBRADO):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(3), Cm(2.8)
    sec.left_margin, sec.right_margin = Cm(3), Cm(2.3)
    sec.header_distance, sec.footer_distance = Cm(0.4), Cm(0.8)

    estilo = doc.styles['Normal']
    estilo.font.name = VISUAL['fonte']
    estilo.element.rPr.rFonts.set(qn('w:eastAsia'), VISUAL['fonte'])
    estilo.font.size = Pt(VISUAL['tamanho'])
    estilo.font.color.rgb = TEXTO

    cab = sec.header.paragraphs[0]
    cab.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if os.path.exists(LOGO):
        cab.add_run().add_picture(LOGO, width=Cm(2.7))
    if os.path.exists(MARCA):
        _imagem_atras_do_texto(cab.add_run(), MARCA, 14.6, 18.7)

    rod = sec.footer.paragraphs[0]
    rod.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _borda_superior(rod)
    e = ESCRITORIO
    for i, linha in enumerate((
        f"{e['telefone']} | {e['instagram']} | {e['site']}",
        f"{e['endereco_timbrado']}  |  E-mail: {e['email']}",
    )):
        r = rod.add_run(linha)
        r.font.name = VISUAL['fonte_rodape']
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0, 0, 0)
        if i == 0:
            r.add_break()
    doc.save(saida)
    return saida


# ============================================================
# ESCRITA
# ============================================================

def novo_documento():
    if not os.path.exists(TIMBRADO):
        montar_timbrado()
    doc = Document(TIMBRADO)
    corpo = doc.element.body
    for el in list(corpo):
        if el.tag != qn('w:sectPr'):
            corpo.remove(el)
    return doc


def _formatar(p, alinhamento=WD_ALIGN_PARAGRAPH.JUSTIFY, recuo=False, depois=6, espaco=None):
    p.alignment = alinhamento
    f = p.paragraph_format
    f.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    f.line_spacing = espaco or VISUAL['espacamento']
    f.space_after = Pt(depois)
    if recuo:
        f.first_line_indent = Cm(VISUAL['recuo_primeira_linha_cm'])
    return p


def titulo(doc, texto, tamanho=14, centralizado=True):
    p = doc.add_paragraph()
    _formatar(p, WD_ALIGN_PARAGRAPH.CENTER if centralizado else WD_ALIGN_PARAGRAPH.JUSTIFY, depois=10)
    r = p.add_run(texto.upper())
    r.bold, r.font.size, r.font.color.rgb = True, Pt(tamanho), DESTAQUE
    return p


def secao(doc, texto):
    p = doc.add_paragraph()
    _formatar(p, depois=4)
    p.paragraph_format.space_before = Pt(10)
    r = p.add_run(texto.upper())
    r.bold, r.font.color.rgb = True, DESTAQUE
    return p


def paragrafo(doc, texto, rotulo=None, recuo=False, negrito=False, tamanho=None, espaco=None):
    p = doc.add_paragraph()
    _formatar(p, recuo=recuo, espaco=espaco)
    if rotulo:
        r = p.add_run(f'{rotulo}: ')
        r.bold, r.font.color.rgb = True, DESTAQUE
        if tamanho:
            r.font.size = Pt(tamanho)
    _texto_com_conferir(p, texto or CONFERIR, negrito, tamanho)
    return p


def _texto_com_conferir(p, texto, negrito=False, tamanho=None):
    """Escreve o texto destacando em vermelho qualquer [CONFERIR] ou [PREENCHER]."""
    for parte in re.split(r'(\[(?:CONFERIR|PREENCHER)[^\]]*\])', str(texto)):
        if not parte:
            continue
        r = p.add_run(parte)
        r.bold = negrito
        if tamanho:
            r.font.size = Pt(tamanho)
        if parte.startswith('[CONFERIR') or parte.startswith('[PREENCHER'):
            r.bold, r.font.color.rgb = True, RGBColor(0xC0, 0x00, 0x00)


def lista(doc, itens, tamanho=None):
    for item in itens:
        p = doc.add_paragraph()
        _formatar(p, depois=2, espaco=1.15)
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.first_line_indent = Cm(-0.4)
        _texto_com_conferir(p, f'• {item}', tamanho=tamanho)


def tabela(doc, cabecalho, linhas, larguras_cm=None, tamanho=9):
    t = doc.add_table(rows=1, cols=len(cabecalho))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, nome in enumerate(cabecalho):
        cel = t.rows[0].cells[i]
        cel.text = ''
        r = cel.paragraphs[0].add_run(nome)
        r.bold, r.font.size = True, Pt(tamanho)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        sombra = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{VISUAL["cor_destaque"]}"/>')
        cel._tc.get_or_add_tcPr().append(sombra)
    for linha in linhas:
        cels = t.add_row().cells
        for i, valor in enumerate(linha):
            cels[i].text = ''
            p = cels[i].paragraphs[0]
            _texto_com_conferir(p, valor if valor not in (None, '') else '-', tamanho=tamanho)
    if larguras_cm:
        # o Word so respeita a largura com autoajuste desligado e a grade da tabela preenchida
        t.autofit = False
        grade = t._tbl.tblGrid
        for i, col in enumerate(grade.findall(qn('w:gridCol'))):
            if i < len(larguras_cm):
                col.set(qn('w:w'), str(int(Cm(larguras_cm[i]).twips)))
        for linha in t.rows:
            for i, w in enumerate(larguras_cm):
                linha.cells[i].width = Cm(w)
    # cabecalho repete no topo de cada pagina e nao fica sozinho no pe da pagina
    trPr = t.rows[0]._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    for cel in t.rows[0].cells:
        for p in cel.paragraphs:
            p.paragraph_format.keep_with_next = True
    doc.add_paragraph()
    return t


# ============================================================
# MODELOS COM {{CAMPOS}}
# ============================================================

CAMPO = re.compile(r'\{\{\s*([A-Z0-9_]+)\s*\}\}')


def _substituir_no_paragrafo(p, dados, faltando):
    if '{{' not in ''.join(r.text for r in p.runs):
        return

    def troca(m):
        chave = m.group(1)
        valor = dados.get(chave)
        if valor in (None, ''):
            faltando.add(chave)
            return f'[PREENCHER {chave}]'
        return str(valor)

    # 1) campo inteiro dentro de um pedaco: troca ali mesmo e preserva a formatacao
    for r in p.runs:
        if '{{' in r.text:
            r.text = CAMPO.sub(troca, r.text)
    # 2) campo quebrado em varios pedacos pelo Word: junta no primeiro pedaco
    texto = ''.join(r.text for r in p.runs)
    if CAMPO.search(texto):
        for r in p.runs[1:]:
            r._r.getparent().remove(r._r)
        p.runs[0].text = CAMPO.sub(troca, texto)


def _todos_paragrafos(doc):
    for p in doc.paragraphs:
        yield p
    for t in doc.tables:
        for linha in t.rows:
            for cel in linha.cells:
                yield from cel.paragraphs
    for sec in doc.sections:
        for parte in (sec.header, sec.footer):
            yield from parte.paragraphs


def preencher_modelo(modelo, dados, saida):
    """Preenche o modelo e salva. Retorna o conjunto de campos que ficaram sem dado."""
    doc = Document(modelo)
    faltando = set()
    for p in _todos_paragrafos(doc):
        _substituir_no_paragrafo(p, dados, faltando)
    doc.save(saida)
    return faltando


# ============================================================
# PDF
# ============================================================

def _tem_word():
    """Microsoft Word instalado (o docx2pdf so funciona com ele: Windows e Mac)."""
    if sys.platform == 'win32':
        return True   # sem Word o docx2pdf so falha e cai no LibreOffice
    if sys.platform == 'darwin':
        return any(os.path.isdir(os.path.join(p, 'Microsoft Word.app'))
                   for p in ('/Applications', os.path.expanduser('~/Applications')))
    return False      # Linux/VPS: so LibreOffice


def caminho_soffice():
    """LibreOffice (soffice) no Windows, no Mac ou no Linux. SOFFICE_PATH no .env tem prioridade."""
    achado = os.getenv('SOFFICE_PATH') or shutil.which('soffice') or shutil.which('libreoffice')
    if achado:
        return achado
    candidatos = [
        os.path.join(os.getenv('ProgramFiles', r'C:\Program Files'), 'LibreOffice', 'program', 'soffice.exe'),
        os.path.join(os.getenv('ProgramFiles(x86)', r'C:\Program Files (x86)'), 'LibreOffice', 'program', 'soffice.exe'),
        '/Applications/LibreOffice.app/Contents/MacOS/soffice',
        os.path.expanduser('~/Applications/LibreOffice.app/Contents/MacOS/soffice'),
        '/opt/homebrew/bin/soffice', '/usr/local/bin/soffice',   # Homebrew (Mac Apple Silicon / Intel)
        '/usr/bin/soffice', '/usr/bin/libreoffice',
    ]
    return next((c for c in candidatos if os.path.exists(c)), None)


def docx_para_pdf(caminho_docx):
    """Converte para PDF ao lado do .docx. Retorna o caminho do PDF ou None.
    Ordem: Word (docx2pdf; Windows e Mac com Word) e, se nao der, LibreOffice.
    PDF_CONVERSOR=libreoffice no .env pula o Word (ex.: Mac sem Word ou rotina sem ninguem na maquina,
    onde o Word do Mac pode ficar esperando a permissao de automacao)."""
    pdf = os.path.splitext(caminho_docx)[0] + '.pdf'
    if os.getenv('PDF_CONVERSOR', '').strip().lower() != 'libreoffice' and _tem_word():
        try:
            from docx2pdf import convert  # usa o Microsoft Word instalado
            convert(caminho_docx, pdf)
            if os.path.exists(pdf):
                return pdf
        except Exception:  # noqa: BLE001 - sem Word ou Word com erro: tenta o LibreOffice
            pass
    soffice = caminho_soffice()
    if soffice:
        import tempfile
        from pathlib import Path
        # perfil proprio: com o LibreOffice aberto na tela, o modo --headless nao converte nada
        perfil = Path(tempfile.gettempdir(), 'caldeira_libreoffice').as_uri()
        try:
            subprocess.run([soffice, f'-env:UserInstallation={perfil}', '--headless', '--convert-to', 'pdf',
                            '--outdir', os.path.dirname(os.path.abspath(caminho_docx)), caminho_docx],
                           capture_output=True, timeout=180)
        except (OSError, subprocess.SubprocessError):
            return None
        if os.path.exists(pdf):
            return pdf
    return None


if __name__ == '__main__':
    print('Timbrado gerado em', montar_timbrado())
