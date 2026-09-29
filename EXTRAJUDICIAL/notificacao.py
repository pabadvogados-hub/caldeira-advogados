"""
Notificacao extrajudicial ao banco/cooperativa - dois modelos do escritorio:

  alongamento  "ASSUNTO: PEDIDO DE ALONGAMENTO DE DIVIDA RURAL" (modelo longo, msg37246)
               I. DOS FATOS / II. DOS DIREITOS / III. DOS PEDIDOS
  contratos    "ASSUNTO: PEDIDO DE CONTRATOS RURAIS" (modelo curto, msg37244) - quando faltam as cedulas

O esqueleto e fixo (qualificacao, tabela das operacoes, transcricao do MCR 2.6.4 e da Sumula 298,
pedidos, fecho, assinaturas). A IA (ia.texto_longo) redige so a apresentacao do produtor, os fatos
e a ligacao dos fatos com as alineas do MCR, a partir do caso.json e dos documentos da pasta.
Sem dado: [CONFERIR ...]/[PREENCHER ...] em vermelho, e isso trava o rascunho no Gmail.
A IA nunca envia nada: o rascunho so e criado com --rascunho-gmail e o advogado clica em Enviar.
"""
import json
import os
import re
import shutil
import sys
import tempfile

sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'NUCLEO'))
import ambiente  # noqa: E402,F401  (carrega .env e caminhos)

from docx import Document  # noqa: E402

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import bancos as base_bancos  # noqa: E402
import escrita as e  # noqa: E402
import registro as reg  # noqa: E402
from docx_caldeira import docx_para_pdf, novo_documento, tabela  # noqa: E402
from config.escritorio import ESCRITORIO, OUTORGADOS  # noqa: E402

# E-mail para o banco responder (o mesmo das notificacoes reais do escritorio). Mudar no .env se precisar.
EMAIL_RESPOSTAS = os.getenv('EMAIL_RESPOSTAS_BANCOS') or 'cfagro.advocacia@gmail.com'
CARENCIA_PADRAO, PARCELAS_PADRAO = 3, 15   # o que a notificacao real pediu (37246); o laudo manda
MARCA_REVISAO = ('MINUTA GERADA PELA IA - PRONTA PARA REVISÃO do Adv. Extrajudicial. '
                 'Esta linha não vai no PDF de envio.')
VALOR =re.compile(r'R\$\s*([\d.]+,\d{2})')
DATA = re.compile(r'\b\d{2}/\d{2}/\d{4}\b')
NUMERO_CEDULA = re.compile(r'(?:n[º°o.]\s*|n[uú]mero\s*|opera[cç][aã]o\s*)([0-9][0-9./-]{4,})', re.I)
NUMEROS = ['zero', 'um', 'dois', 'três', 'quatro', 'cinco', 'seis', 'sete', 'oito', 'nove', 'dez', 'onze',
           'doze', 'treze', 'quatorze', 'quinze', 'dezesseis', 'dezessete', 'dezoito', 'dezenove', 'vinte']


# ============================================================
# DADOS
# ============================================================

credor_nao_bancario = reg.credor_nao_bancario


def _v(valor, campo):
    return valor if valor not in (None, '') else f'[CONFERIR {campo}]'


def _moeda(v):
    return 'R$ ' + f'{v:,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.')


def _valor(txt):
    m = VALOR.search(txt or '')
    return float(m.group(1).replace('.', '').replace(',', '.')) if m else None


def _por_extenso(n):
    return f'{n:02d} ({NUMEROS[n]})' if 0 <= n < len(NUMEROS) else str(n)


def _genero(q):
    ec = (q.get('estado_civil') or '').strip().lower()
    return 'f' if ec.endswith('a') else 'm'


def endereco(q):
    numero = q.get('numero') or ''
    numero = f'n° {numero}' if numero[:1].isdigit() else numero   # "Lote 32" fica como esta
    partes = [q.get('logradouro'), numero, q.get('bairro')]
    txt = ', '.join(x for x in partes if x)
    if q.get('cep'):
        txt += f", CEP: {q['cep']}"
    if q.get('cidade'):
        txt += f", município de {q['cidade']}" + (f"/{q['uf']}" if q.get('uf') else '')
    return txt or '[CONFERIR endereço]'


def qualificacao_notificante(caso):
    q = caso.get('qualificacao') or {}
    f = _genero(q) == 'f'
    nac = (q.get('nacionalidade') or '').lower()
    if nac.endswith('a') and not f:
        nac = nac[:-1] + 'o'
    rg = ''
    if q.get('rg'):
        rg = (f"{'portadora' if f else 'portador'} da Carteira de Identidade RG n. {q['rg']}"
              f"{' ' + q['orgao_emissor'] if q.get('orgao_emissor') else ''} e ")
    texto = (f", {_v(nac, 'nacionalidade')}, {_v((q.get('estado_civil') or '').lower(), 'estado civil')}, "
             f"{_v((q.get('profissao') or '').lower(), 'profissão')}, maior e capaz, {rg}"
             f"devidamente {'inscrita' if f else 'inscrito'} no CPF sob o n. {_v(q.get('cpf'), 'CPF')}, "
             f"{'residente e domiciliada' if f else 'residente e domiciliado'} na {endereco(q)}, ")
    return [('NOTIFICANTE: ', 'o'), (reg.nome_cliente(caso).upper(), 'o'), (texto, 'n'),
            ('representado por seus advogados devidamente constituídos e qualificados nos termos da '
             'procuração em anexo.', 'b')]


def qualificacao_notificado(banco, entrada):
    if not entrada:
        return [('NOTIFICADO: ', 'o'), (banco.upper(), 'o'),
                (', [CONFERIR razão social, CNPJ e endereço da agência: banco fora de config/bancos_emails.json].', 'n')]
    nome = (entrada.get('razao_social') or '').upper() or f"{banco.upper()} [CONFERIR razão social]"
    agencia = f", agência {entrada['agencia']}" if entrada.get('agencia') else ''
    local = entrada.get('endereco') or '[CONFERIR endereço da agência]'
    municipio = f", município de {entrada['municipio']}" if entrada.get('municipio') else ''
    return [('NOTIFICADO: ', 'o'), (nome, 'o'),
            (f", CNPJ: {_v(entrada.get('cnpj'), 'CNPJ')}{agencia}, com endereço na {local}{municipio}.", 'n')]


def linhas_operacoes(ops, texto_cedulas):
    """Linhas da tabela de operacoes. Valor dito na reuniao e nao achado na cedula leva [CONFERIR]."""
    linhas, total, confirmado = [], 0.0, bool(ops)
    for o in ops:
        m = NUMERO_CEDULA.search(f"{o.get('instrumento', '')} {o.get('observacao', '')}")
        numero = m.group(1) if m else '[CONFERIR nº da cédula]'
        instr = o.get('instrumento') or ''
        if re.search(r'n[aã]o informad', instr, re.I):
            instr = '[CONFERIR tipo do título]'
        fin = o.get('finalidade') or ''
        if re.search(r'n[aã]o informad', fin, re.I):
            fin = ''
        titulo_fin = ' - '.join(x for x in (fin, instr) if x) or '[CONFERIR título e finalidade]'
        v = _valor(o.get('valor'))
        if v is None:
            valor_txt, confirmado = '[CONFERIR valor]', False
        else:
            total += v
            valor_txt = _moeda(v)
            achou = VALOR.search(o.get('valor')).group(1) in (texto_cedulas or '')
            if not achou or re.search(r'aproximad', o.get('valor') or '', re.I):
                valor_txt += ' [CONFERIR na cédula]'
                confirmado = False
        datas = DATA.findall(o.get('vencimentos') or '')
        venc = ', '.join(datas) if datas else '[CONFERIR vencimento]'
        s = (o.get('situacao') or '').lower()
        if 'execu' in s:
            sit = 'Em execução'
        elif 'a vencer' in s:
            sit = 'A vencer'
        elif 'vencid' in s or 'inadimpl' in s:
            sit = 'Vencida'
        elif 'renegoc' in s or 'prorrog' in s:
            sit = 'Renegociada'
        elif 'em dia' in s:
            sit = 'Em dia'
        else:
            sit = '[CONFERIR situação]'
        linhas.append([numero, titulo_fin, valor_txt, venc, sit])
    return linhas, (total if total else None), confirmado


# ============================================================
# IA
# ============================================================

def _sistema():
    e_ = ESCRITORIO
    return f"""Voce e advogado(a) do {e_['nome']} ({e_['cidade']}/{e_['uf']}), escritorio que defende o produtor rural contra bancos e cooperativas de credito. Voce redige blocos da NOTIFICACAO EXTRAJUDICIAL "PEDIDO DE ALONGAMENTO DE DIVIDA RURAL", no padrao das notificacoes reais do escritorio (secao 5.2 do material de referencia e a peca 37246).

O sistema monta o esqueleto fixo: qualificacao das partes, tabela das operacoes com o valor total, transcricao do MCR 2.6.4 e da Sumula 298/STJ, pedidos (alongamento e carencia, retirada do debito automatico, envio das cedulas, comunicacao exclusiva com os advogados), fecho e assinaturas. Voce redige SOMENTE os blocos pedidos.

Regras inegociaveis:
1. Nunca invente dado: nome, numero, valor, area, data, percentual, numero de decreto, preco, quantidade de animais ou sacas. Use somente o que estiver em <dados_do_caso>, <cedulas> e <laudos>. Fato publico de Rondonia (decretos, precos da arroba, notas tecnicas) so se estiver no material de referencia E combinar com a atividade (pecuaria x lavoura) e o periodo do caso; onde o material marca [CONFERIR], mantenha a marca.
2. Faltou o dado: escreva [PREENCHER descricao curta do que falta] ou [CONFERIR descricao] no ponto exato. Valor ou data dito "aproximadamente" pelo produtor: escreva "aproximadamente" e acrescente [CONFERIR].
3. A notificacao vai para o banco: nao mencione reuniao, Closer, transcricao, honorarios, estrategia, riscos, fraquezas do caso (ex.: "nao tem laudo"), ameacas do gerente, nem dados pessoais (CPF, RG, telefone). Nada de jurisprudencia: os fundamentos legais ja estao no esqueleto.
4. Tom das pecas do escritorio: formal, enfatico, protetivo do produtor, terceira pessoa ("o Notificante"; o banco e "o Notificado" ou "a Instituicao Financeira"). Nunca "Vossa Excelencia" (nao e peticao).
5. Sem laudo em <laudos>: a secao "4) DO LAUDO AGRARIO" e a "5) DA ESTIMATIVA DE PAGAMENTO" ficam com uma unica linha [PREENCHER ...] explicando o que o laudo precisa trazer. Com laudo: resuma os numeros do laudo (periodos de deficit hidrico, producao projetada x obtida, preco, faturamento, ano de retomada, numero de parcelas).
6. Formatacao do texto (o sistema converte): "## 1) TITULO" para subtitulo; "### a) Titulo" para alinea; "- Titulo: texto" para marcador; paragrafos simples em linhas proprias; **trecho** para destaque. Sem tabelas, sem outras marcacoes.

Responda EXATAMENTE neste formato, com os marcadores:
<<<APRESENTACAO>>>
(1 ou 2 paragrafos: o Notificante, a atividade, o municipio, a area e a relacao com a instituicao; ate 180 palavras)
<<<FATOS>>>
(vem logo depois da tabela das operacoes. Comece com 1 paragrafo dizendo que a dificuldade nao decorre de ma gestao e com marcadores dos fatores estruturais que se aplicam ao caso. Depois as secoes, nesta ordem, pulando a que nao tiver base: "## 1) DAS CIRCUNSTANCIAS NATURAIS ADVERSAS" (com alineas a, b...), "## 2) DAS CIRCUNSTANCIAS MERCADOLOGICAS ADVERSAS", "## 3) DAS PRINCIPAIS CAUSAS ADVERSAS AO AGRONEGOCIO" (cada causa com a alinea do MCR 2.6.4 entre parenteses, so as que o caso sustenta), "## 4) DO LAUDO AGRARIO", "## 5) DA ESTIMATIVA DE PAGAMENTO". De 600 a 1600 palavras.)
<<<FUNDAMENTOS>>>
(1 a 3 paragrafos ligando os fatos do caso as alineas do MCR 2.6.4 que se aplicam; vai depois da transcricao do MCR e de um paragrafo fixo que ja diz que atestar a necessidade e apurar a capacidade de pagamento sao deveres da instituicao - nao repita isso; ate 250 palavras)
<<<PARAMETROS>>>
CARENCIA_ANOS: (numero inteiro so se o laudo trouxer; senao vazio)
PARCELAS_ANUAIS: (numero inteiro so se o laudo trouxer; senao vazio)
<<<AVISOS>>>
(marcadores para o advogado revisor: o que conferir antes de enviar; nao vai para o banco)
"""


def redigir(caso, banco, ops, cedulas, laudos, anterior=None):
    import ia
    t = caso.get('triagem') or {}
    q = caso.get('qualificacao') or {}
    dados = {
        'notificante': {'nome': reg.nome_cliente(caso), 'profissao': q.get('profissao'),
                        'municipio': q.get('cidade'), 'uf': q.get('uf')},
        'credor_notificado': banco,
        'operacoes_com_este_credor': ops,
        'outros_credores': [b for b in reg.bancos_do_caso(caso) if not reg.mesmo_banco(b, banco)],
        'atividade_rural': t.get('atividade_rural'),
        'resumo_do_caso': t.get('resumo_caso'),
        'linha_do_tempo': t.get('linha_do_tempo'),
    }
    conteudo = f'<dados_do_caso>\n{json.dumps(dados, ensure_ascii=False, indent=1)}\n</dados_do_caso>'
    conteudo += f"\n\n<cedulas>\n{cedulas or '(nenhuma cedula na pasta 03 ainda)'}\n</cedulas>"
    conteudo += f"\n\n<laudos>\n{laudos or '(nenhum laudo na pasta 06 ainda)'}\n</laudos>"
    if anterior:
        conteudo += (f"\n\n<notificacao_anterior>Notificacao enviada em {reg.br(anterior.get('enviada_em'))}, "
                     f"situacao: {reg.situacao(anterior)}. Esta e a reiteracao.</notificacao_anterior>")
    conteudo += f'\n\nRedija os blocos da notificacao ao {banco}.'
    sistema = _sistema()
    dna = ambiente.ler_base('DNA_PECAS.md')
    if dna:
        sistema += f'\n<material_de_referencia>\n{dna}\n</material_de_referencia>\n'
    bruto = ia.texto_longo(sistema, conteudo, max_tokens=20000)
    return _separar(bruto)


def _separar(bruto):
    blocos = {'APRESENTACAO': '', 'FATOS': '', 'FUNDAMENTOS': '', 'PARAMETROS': '', 'AVISOS': ''}
    partes = re.split(r'<<<([A-Z]+)>>>', bruto or '')
    if len(partes) < 3:
        blocos['FATOS'] = bruto or ''
    for i in range(1, len(partes) - 1, 2):
        if partes[i] in blocos:
            blocos[partes[i]] = partes[i + 1].strip()
    par = {}
    for chave in ('CARENCIA_ANOS', 'PARCELAS_ANUAIS'):
        m = re.search(chave + r'\s*:\s*(\d+)', blocos['PARAMETROS'])
        par[chave] = int(m.group(1)) if m else None
    blocos['parametros'] = par
    blocos['avisos'] = [a.lstrip('-•* ').replace('**', '').strip() for a in blocos['AVISOS'].splitlines()
                        if a.strip()]
    return blocos


def sem_ia(caso, banco):
    """Texto-base quando nao ha IA (sem chave ou --sem-ia): o advogado completa os [PREENCHER]."""
    t = caso.get('triagem') or {}
    ar = t.get('atividade_rural') or {}
    q = caso.get('qualificacao') or {}
    municipio = ar.get('municipio_propriedade') or q.get('cidade') or '[PREENCHER município]'
    atividade = ar.get('culturas_rebanho') or '[PREENCHER atividade]'
    apres = (f"O NOTIFICANTE é produtor rural em {municipio}, dedicado a {atividade.lower()}"
             f"{', em área de ' + ar['area_propriedade'] + ' [CONFERIR]' if ar.get('area_propriedade') else ''}, "
             f"e mantém operações de crédito rural com o {banco.upper()}, utilizadas exclusivamente no custeio "
             'e no investimento da sua atividade produtiva. [PREENCHER histórico do produtor com a instituição]')
    fatos = '\n'.join([
        'A inviabilidade de honrar os compromissos nas datas originais não decorre de má gestão, mas de fatores '
        'adversos de ordem climática e de mercado que atingiram a atividade do Notificante.',
        '## 1) DAS CIRCUNSTÂNCIAS NATURAIS ADVERSAS',
        f"[PREENCHER eventos adversos com meses e anos: {ar.get('causa_da_perda') or 'causa da perda'}]",
        '## 2) DAS CIRCUNSTÂNCIAS MERCADOLÓGICAS ADVERSAS',
        '[PREENCHER preços obtidos x esperados, se houver]',
        '## 3) DAS PRINCIPAIS CAUSAS ADVERSAS AO AGRONEGÓCIO',
        '[PREENCHER causas ligadas às alíneas "a", "b", "c" e "d" do MCR 2.6.4]',
        '## 4) DO LAUDO AGRÁRIO',
        '[PREENCHER resumo do Laudo de Frustração de Safra: déficit hídrico por período, produção projetada x '
        'obtida, preço e faturamento]',
        '## 5) DA ESTIMATIVA DE PAGAMENTO',
        '[PREENCHER ano de retomada da solvência e número de parcelas anuais, conforme o laudo de capacidade de pagamento]',
    ])
    return {'APRESENTACAO': apres, 'FATOS': fatos,
            'FUNDAMENTOS': '[PREENCHER ligação dos fatos do caso com as alíneas do MCR 2.6.4]',
            'parametros': {'CARENCIA_ANOS': None, 'PARCELAS_ANUAIS': None},
            'avisos': ['Texto sem IA: completar todos os [PREENCHER].']}


# ============================================================
# DOCUMENTO
# ============================================================

def _cabecalho(doc, caso, banco, entrada, tipo, anterior):
    e.titulo_notificacao(doc)
    e.pedacos(doc, qualificacao_notificante(caso), depois=10)
    e.pedacos(doc, qualificacao_notificado(banco, entrada), depois=10)
    assunto = reg.TIPOS[tipo] + (' - REITERAÇÃO' if anterior else '')
    e.pedacos(doc, [('ASSUNTO: ', 'o'), (assunto, 'n')], depois=10)
    e.pedacos(doc, [('PREZADOS SENHORES,', 'o')], depois=10)
    if anterior:
        resp = ('que até a presente data não obteve resposta' if reg.situacao(anterior) in
                ('SEM RESPOSTA', 'PRAZO VENCIDO', 'AGUARDANDO RESPOSTA') else 'cujo retorno não atendeu ao pedido')
        e.pedacos(doc, [('O Notificante ', 'n'), ('reitera', 'o'),
                        (f" a notificação extrajudicial encaminhada a esta instituição em "
                         f"{reg.br(anterior.get('enviada_em')) if anterior.get('enviada_em') else '[CONFERIR data]'}, "
                         f'{resp}, e renova integralmente os pedidos abaixo.', 'n')], recuo=True)


def _fecho(doc):
    doc.add_paragraph()
    e.pedacos(doc, [(f"{ESCRITORIO['cidade']} - {ESCRITORIO['uf']}, {reg.extenso()}.", 'n')],
              alinhamento=1, depois=12)
    e.assinaturas(doc, OUTORGADOS)
    e.nota_revisao(doc, MARCA_REVISAO)


def montar_contratos(caso, banco, entrada, anterior):
    doc = novo_documento()
    _cabecalho(doc, caso, banco, entrada, 'contratos', anterior)
    e.pedacos(doc, [('Pelo presente, o Notificante, devidamente representado por seus advogados, conforme '
                     'procuração anexa, vem, respeitosamente, requerer o seguinte:', 'n')], recuo=True)
    e.pedacos(doc, [('A) ', 'n'), ('O fornecimento de ', 'n'),
                    ('cópias integrais de todos os Contratos Rurais firmados entre o Notificante e o Notificado', 'o'),
                    (', incluindo:', 'n')], antes=6)
    for item in ('Cédulas de Crédito Rural (simples ou vinculadas);',
                 'Contratos de Financiamento para Custeio Agrícola e Pecuário;',
                 'Contratos de Investimento Rural;',
                 'Contratos de Comercialização ou quaisquer outros instrumentos contratuais relacionados às '
                 'operações de crédito rural.'):
        e.marcador(doc, item)
    e.pedacos(doc, [('Tal solicitação se dá em razão da necessidade de análise documental e acompanhamento das '
                     'condições pactuadas.', 'n')], recuo=True, antes=6)
    e.pedacos(doc, [('B) ', 'n'), ('Solicita-se', 'o'),
                    (', ainda, caso existam aditivos contratuais, termos de prorrogação, renegociação ou '
                     'repactuação, que os mesmos também sejam disponibilizados.', 'n')], antes=6)
    e.pedacos(doc, [('As informações poderão ser encaminhadas por meio digital, preferencialmente em formato PDF, '
                     f'para o e-mail: {EMAIL_RESPOSTAS}.', 'n')], recuo=True, antes=6)
    _fecho(doc)
    return doc, {'avisos': []}


def montar_alongamento(caso, banco, entrada, anterior, textos, cedulas, carencia=None, parcelas=None):
    ops = reg.operacoes_do_banco(caso, banco)
    doc = novo_documento()
    _cabecalho(doc, caso, banco, entrada, 'alongamento', anterior)
    e.pedacos(doc, [('O NOTIFICANTE, produtor rural, vem, respeitosamente, por meio da presente ', 'n'),
                    ('NOTIFICAÇÃO EXTRAJUDICIAL', 'o'), (', expor e requerer o seguinte:', 'n')], recuo=True)

    e.secao_romana(doc, 'I. DOS FATOS:')
    e.markdown(doc, textos.get('APRESENTACAO'))
    linhas, total, confirmado = linhas_operacoes(ops, cedulas)
    total_txt = _moeda(total) if total else '[CONFERIR valor total]'
    if total and not confirmado:
        total_txt += ' [CONFERIR valor atual com a cédula/extrato]'
    e.pedacos(doc, [('Apesar de sua gestão responsável, o Notificante se vê compelido a solicitar a renegociação '
                     'dos seguintes Instrumentos de Crédito Rural, celebrados junto ao ', 'n'),
                    (banco.upper(), 'o'), (', que totalizam um valor atual de ', 'n'), (total_txt, 'o'), (':', 'n')],
              recuo=True)
    if linhas:
        tabela(doc, ['Nº da operação', 'Título e finalidade', 'Valor', 'Vencimento', 'Situação'], linhas,
               [3.0, 5.2, 3.4, 2.6, 2.2], tamanho=9)
    else:
        e.pedacos(doc, [('[PREENCHER tabela das operações com este banco: nº, título, valor, vencimento e situação]', 'n')])
    e.markdown(doc, textos.get('FATOS'))

    e.secao_romana(doc, 'II. DOS DIREITOS:')
    e.pedacos(doc, [('O Manual de Crédito Rural (MCR), no capítulo 2, seção 6, item 4, autoriza a instituição '
                     'financeira a prorrogar a dívida rural, aos mesmos encargos financeiros pactuados nos instrumentos '
                     'de crédito, desde que comprovada a dificuldade temporária de reembolso do mutuário:', 'n')],
              recuo=True)
    e.citacao(doc, '“4 - Fica a instituição financeira autorizada a prorrogar a dívida, aos mesmos encargos financeiros '
                   'pactuados no instrumento de crédito, desde que o mutuário comprove a dificuldade temporária para '
                   'reembolso do crédito em razão de uma ou mais entre as situações abaixo, e que a instituição financeira '
                   'ateste a necessidade de prorrogação e demonstre a capacidade de pagamento do mutuário: '
                   '(Res CMN 4.883 art 1º; Res CMN 4.905 art 1º)')
    e.citacao(doc, 'a) dificuldade de comercialização dos produtos; (Res CMN 4.883 art 1º);')
    e.citacao(doc, 'b) frustração de safras, por fatores adversos; (Res CMN 4.883 art 1º);')
    e.citacao(doc, 'c) eventuais ocorrências prejudiciais ao desenvolvimento das explorações. (Res CMN 4.883 art 1º);')
    e.citacao(doc, 'd) dificuldades no fluxo de caixa do mutuário, devido ao impacto acumulado de perdas de safra '
                   'decorrentes de eventos climáticos adversos em safras anteriores, que gerem aumento do endividamento '
                   'no Sistema Nacional de Crédito Rural - SNCR e impossibilitem o reembolso integral das operações de '
                   'crédito rural. (Res CMN 5.229 art 5º)”')
    e.pedacos(doc, [('Registre-se que o próprio item 2.6.4 exige do mutuário apenas a comprovação da dificuldade '
                     'temporária; atestar a necessidade da prorrogação e apurar a capacidade de pagamento são deveres '
                     'da instituição financeira.', 'n')], recuo=True)
    e.markdown(doc, textos.get('FUNDAMENTOS'))
    e.pedacos(doc, [('Em complemento ao MCR 2.6.4, a Súmula 298 do Superior Tribunal de Justiça (STJ) estabelece que o '
                     'alongamento das dívidas oriundas de crédito rural não constitui mera faculdade da instituição '
                     'financeira, mas sim um ', 'n'), ('direito do devedor', 'o'), (':', 'n')], recuo=True)
    e.citacao(doc, '“O alongamento de dívida originada de crédito rural não constitui faculdade da instituição '
                   'financeira, mas, direito do devedor nos termos da lei.”')
    e.pedacos(doc, [('Nesse sentido, ao preencher os requisitos estabelecidos pela legislação e pelos normativos '
                     'aplicáveis, o produtor rural tem o direito de obter a repactuação do débito, cabendo à instituição '
                     'financeira o dever de viabilizar a renegociação nos termos fixados. A recusa injustificada pode '
                     'configurar descumprimento da norma, sujeitando-a à revisão judicial da negativa.', 'n')], recuo=True)
    e.pedacos(doc, [('O direito ao alongamento do crédito rural está também respaldado pelo Código de Defesa do '
                     'Consumidor (CDC), quando aplicável, e pela legislação civil, que prevê a revisão contratual em '
                     'casos de onerosidade excessiva e imprevisibilidade, conforme os artigos 317 e 478 do Código Civil. '
                     'Ademais, a Lei nº 4.829/1965, que institui o crédito rural, ', 'n'),
                    ('reforça a função social do financiamento agrícola, garantindo condições justas e equilibradas '
                     'para o setor.', 'o')], recuo=True)
    e.pedacos(doc, [('Assim, diante da situação produtiva e financeira demonstrada, o Notificante se enquadra nos '
                     'requisitos do MCR e nas disposições legais aplicáveis, sendo titular do direito ao Alongamento de '
                     'suas Dívidas Rurais junto ao Notificado, com a manutenção das condições originais de encargos '
                     'financeiros pactuados, nos termos da regulamentação vigente.', 'n')], recuo=True)

    e.secao_romana(doc, 'III. DOS PEDIDOS:')
    e.pedacos(doc, [('Pelo presente, o Notificante, devidamente representado por seus advogados, conforme procuração '
                     'anexa, vem, respeitosamente, requerer:', 'n')], recuo=True)
    par = textos.get('parametros') or {}
    c = carencia or par.get('CARENCIA_ANOS')
    n = parcelas or par.get('PARCELAS_ANUAIS')
    conferir = '' if (c and n) else ' [CONFERIR carência e parcelas com o laudo de capacidade de pagamento]'
    c, n = c or CARENCIA_PADRAO, n or PARCELAS_PADRAO
    e.item_numerado(doc, 1, 'Alongamento e Carência',
                    'Solicita-se a prorrogação dos prazos de pagamento das operações acima relacionadas, com as mesmas '
                    f'condições iniciais pactuadas, com um período de carência de **{_por_extenso(c)} anos** e o '
                    f'parcelamento em **{_por_extenso(n)} parcelas anuais**, nos termos das normativas aplicáveis.'
                    + conferir)
    e.item_numerado(doc, 2, 'Retirada do Débito Automático',
                    'Imediata suspensão da modalidade de débito automático incidente sobre quaisquer valores '
                    'vinculados aos contratos de crédito rural firmados entre as partes.')
    e.item_numerado(doc, 3, 'Envio de cédulas e comprovação de adimplência/pagamentos',
                    'Requer-se que a Instituição Financeira encaminhe cópia integral e legível de todas as Cédulas de '
                    'Crédito Rural (CCR) e demais instrumentos contratuais vigentes firmados entre as partes (incluindo '
                    'eventuais aditivos, renegociações e termos correlatos), bem como informações e documentos que '
                    'comprovem a situação de cada contrato;')
    e.item_numerado(doc, 4, 'Comunicação Exclusiva com os Advogados',
                    'O Notificante estabelece que toda e qualquer comunicação referente ao presente pedido deverá ser '
                    'realizada exclusivamente por intermédio de seus advogados constituídos, através do e-mail '
                    f"{EMAIL_RESPOSTAS} ou pelo contato telefônico {ESCRITORIO['telefone']}.")
    e.pedacos(doc, [('Ressalta-se que qualquer contato direto da agência ou de seus representantes com o Notificante '
                     'ensejará a aplicação de multa no valor correspondente a ', 'n'),
                    ('10 (dez) salários mínimos', 'o'), ('.', 'n')], recuo=True, antes=6)
    _fecho(doc)
    return doc, {'avisos': textos.get('avisos') or [], 'carencia': c, 'parcelas': n}


def _nome_arquivo(base, caso, banco):
    pasta = reg.pasta_extra(base)
    raiz = (f'{reg.nome_arquivo_cliente(caso)} - Notificacao Extrajudicial - '
            f'{reg.arquivo_seguro(banco).upper()} - {reg.hoje().strftime("%d-%m-%Y")}')
    caminho = os.path.join(pasta, raiz + '.docx')
    usados = {os.path.abspath(n.get('arquivo') or '') for n in reg.extrajudicial(caso)['notificacoes']}
    i = 2
    while os.path.abspath(caminho) in usados:
        caminho = os.path.join(pasta, f'{raiz} ({i}).docx')
        i += 1
    return caminho


def gerar(base, caso, banco, tipo='alongamento', usar_ia=True, carencia=None, parcelas=None):
    """Gera DOCX + PDF da notificacao e registra no caso.json. Retorna a notificacao registrada."""
    q = caso.get('qualificacao') or {}
    municipio = q.get('cidade') or ((caso.get('triagem') or {}).get('atividade_rural') or {}).get('municipio_propriedade', '')
    entrada = base_bancos.achar(banco, municipio)
    ext = reg.extrajudicial(caso)
    ultima = reg.ultima(caso, banco)
    if ultima and not ultima.get('enviada_em'):
        ext['notificacoes'].remove(ultima)   # minuta anterior nao enviada: substitui
        print(f'   {banco}: minuta anterior (não enviada) substituída.')
    anterior = next((x for x in reversed(reg.notificacoes_do_banco(caso, banco)) if x.get('enviada_em')), None)

    if tipo == 'contratos':
        doc, info = montar_contratos(caso, banco, entrada, anterior)
    else:
        cedulas = reg.texto_documentos(base, 'contratos', banco)
        laudos = reg.texto_documentos(base, 'frustracao', limite_total=60000)
        if not cedulas:
            print(f'   {banco}: nenhuma cédula na pasta 03 (valores e nº da cédula ficam [CONFERIR]; '
                  'se o cliente não tem as cédulas, considere --tipo contratos).')
        if usar_ia and ambiente.tem_credencial('ANTHROPIC_API_KEY'):
            print(f'   {banco}: IA redigindo os fatos (pode levar alguns minutos)...')
            textos = redigir(caso, banco, reg.operacoes_do_banco(caso, banco), cedulas, laudos, anterior)
        else:
            print(f'   {banco}: sem IA (sem ANTHROPIC_API_KEY ou --sem-ia): texto-base com [PREENCHER].')
            textos = sem_ia(caso, banco)
        doc, info = montar_alongamento(caso, banco, entrada, anterior, textos, cedulas, carencia, parcelas)

    caminho = _nome_arquivo(base, caso, banco)
    doc.save(caminho)
    pdf = docx_para_pdf(caminho)
    pend = reg.pendencias_docx(caminho)
    n = reg.nova_notificacao(caso, banco, tipo, caminho, pdf, pend)
    n['avisos_ia'] = info.get('avisos') or []
    if info.get('carencia'):
        n['pedido'] = {'carencia_anos': info['carencia'], 'parcelas_anuais': info['parcelas']}
    n['email_destino'] = (entrada or {}).get('email') or None
    n['proxima_acao'] = reg.proxima_acao_banco(caso, banco)
    return n


# ============================================================
# RASCUNHO NO GMAIL (nunca envia)
# ============================================================

def pdf_de_envio(caminho_docx):
    """PDF do .docx ATUAL (pode ter sido editado no Word) sem a linha de revisao da minuta."""
    doc = Document(caminho_docx)
    for p in list(doc.paragraphs):
        if 'PRONTA PARA REVISÃO' in p.text:
            p._p.getparent().remove(p._p)
    tmp = tempfile.mkdtemp(prefix='notif_')
    try:
        copia = os.path.join(tmp, os.path.basename(caminho_docx))
        doc.save(copia)
        pdf_tmp = docx_para_pdf(copia)
        if not pdf_tmp:
            return None
        destino = os.path.splitext(caminho_docx)[0] + '.pdf'
        shutil.copyfile(pdf_tmp, destino)
        return destino
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def _corpo_email(caso, n, com_procuracao):
    pedido = ('pedido de alongamento (prorrogação) da dívida rural' if n['tipo'] == 'alongamento'
              else 'pedido de cópia integral dos contratos rurais')
    assinam = '\n'.join(f"{a['nome']} - {a['oab']}" for a in OUTORGADOS)
    return (f"Prezados Senhores,\n\n"
            f"Na qualidade de advogados de {reg.nome_cliente(caso).upper()}, encaminhamos em anexo NOTIFICAÇÃO "
            f"EXTRAJUDICIAL referente às operações de crédito rural mantidas com esta instituição, com {pedido}"
            f"{', acompanhada da procuração' if com_procuracao else ''}.\n\n"
            f"Solicitamos a confirmação do recebimento desta mensagem e que toda comunicação sobre o assunto seja feita "
            f"exclusivamente com este escritório, pelo e-mail {EMAIL_RESPOSTAS} ou pelo telefone "
            f"{ESCRITORIO['telefone']}.\n\nAtenciosamente,\n\n{assinam}\n{ESCRITORIO['nome']}\n"
            f"{ESCRITORIO['telefone']} | {ESCRITORIO['site']}\n")


def criar_rascunho(base, caso, n):
    """Cria o rascunho no Gmail do escritorio. Travas: marcas no .docx, e-mail do banco, credencial. NAO ENVIA."""
    banco = n['banco']
    pend = reg.pendencias_docx(n['arquivo'])
    n['pendencias'] = pend
    if pend:
        print(f'   {banco}: rascunho TRAVADO. Resolver no .docx ({len(pend)} marca(s)):')
        for p in pend[:12]:
            print(f'     - {p}')
        return False
    if n.get('enviada_em'):
        print(f'   {banco}: esta notificação já foi registrada como enviada; rascunho não criado.')
        return False
    q = caso.get('qualificacao') or {}
    entrada = base_bancos.achar(banco, q.get('cidade', ''))
    if not entrada or not base_bancos.emails_validos(entrada.get('email')):
        print(f'   {banco}: rascunho TRAVADO. Sem e-mail conferido em config/bancos_emails.json.')
        return False
    if entrada.get('_ambiguo'):
        print(f"   {banco}: rascunho TRAVADO. {entrada['_ambiguo']} (preencher o município da agência na base).")
        return False
    import gmail_integration
    if not gmail_integration.configurado():
        print(f'   {banco}: Gmail sem credencial (config/credentials_gmail.json). Modo seguro: só os arquivos.')
        return False
    pdf = pdf_de_envio(n['arquivo'])
    if not pdf:
        print(f'   {banco}: não consegui gerar o PDF (instalar Word ou LibreOffice).')
        return False
    anexos = [pdf]
    procuracao = reg.procuracao_assinada(base)
    if procuracao:
        anexos.append(procuracao)
    else:
        print(f'   {banco}: AVISO procuração assinada não encontrada em 00 CONTRATACAO; anexe no Gmail antes de enviar.')
    if n['tipo'] == 'alongamento':
        anexos += [a for a in reg.arquivos_item(base, 'frustracao') if a.lower().endswith('.pdf')]
    assunto = (f"NOTIFICAÇÃO EXTRAJUDICIAL - {reg.TIPOS[n['tipo']]}"
               f"{' - REITERAÇÃO' if n.get('numero', 1) > 1 else ''} - {reg.nome_cliente(caso).upper()}")
    try:
        r = gmail_integration.criar_rascunho(entrada['email'], assunto, _corpo_email(caso, n, bool(procuracao)),
                                             anexos, cc=entrada.get('email_copia'),
                                             nome_remetente=ESCRITORIO['nome'])
    except Exception as erro:
        print(f'   {banco}: ERRO ao criar o rascunho no Gmail: {erro}')
        return False
    n.update({'rascunho_id': r['id'], 'rascunho_em': reg.iso(reg.hoje()), 'email_destino': entrada['email'],
              'anexos': [os.path.basename(a) for a in anexos], 'pdf': pdf})
    reg.evento(n, f"Rascunho criado no Gmail para {entrada['email']} (id {r['id']})")
    print(f"   {banco}: RASCUNHO criado no Gmail para {entrada['email']} com {len(anexos)} anexo(s). "
          'Revise e clique em Enviar; depois rode registrar-envio.')
    return True
